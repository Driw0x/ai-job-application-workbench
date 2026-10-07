import shutil
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from docx import Document


REPO_ROOT = Path(__file__).resolve().parents[1]
TEMP_ARTIFACT_PARENTS = (REPO_ROOT, REPO_ROOT / "backend")
TEMP_ARTIFACT_PATHS = (
    REPO_ROOT / ".tmp" / "pytest-company",
    REPO_ROOT / "backend" / ".pytest-tmp-cover-letter",
    REPO_ROOT / "backend" / ".test-tmp",
)


def _temporary_artifacts() -> set[Path]:
    artifacts = {
        path.resolve()
        for parent in TEMP_ARTIFACT_PARENTS if parent.is_dir()
        for path in parent.iterdir()
        if path.is_dir()
        and (path.name == "pytest-job-tracker-run" or path.name.startswith("pytest-cache-files-"))
    }
    return artifacts | {path.resolve() for path in TEMP_ARTIFACT_PATHS if path.exists()}


def _test_repository(repo: Path) -> None:
    career = repo / "01_Career"
    career.mkdir(parents=True)
    (career / "Profil.md").write_text(
        "# Profil\n\n**Nom** : Alex Example\n**Localisation** : Example City\n"
        "**Adresse** : 1 rue des Tests\n**Code postal** : 12345\n**Ville** : Example City\n"
        "**Email** : alex@example.test\n**Téléphone** : 01 23 45 67 89\n\nMaster 2 Computer Science.\n",
        encoding="utf-8",
    )
    for name in ("Domaines.md", "Stage M2.md"):
        (career / name).write_text("# Intelligence artificielle\n\nStage M2 Computer Science.\n", encoding="utf-8")

    titles = ("Search Prototype", "Document Workflow", "Model Experiment", "Planning Prototype", "Image Classifier")
    for title in titles:
        project = repo / "02_Projects" / title / f"{title}.md"
        project.parent.mkdir(parents=True)
        project.write_text(f"# {title}\n\nProjet de test documenté.\n", encoding="utf-8")

    template = career / "CV/Templates/resume-template.docx"
    template.parent.mkdir(parents=True)
    from docx.shared import Cm, Pt

    document = Document()
    document.styles["Normal"].font.size = Pt(8.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(0)
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1)
    section.left_margin = section.right_margin = Cm(1)
    for text in ("Alex Example", "alex@example.test", "Professional profile"):
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    left.width = Cm(6)
    right.width = Cm(13)
    left.paragraphs[0].text = "PROFIL"
    left.add_paragraph("Documented candidate profile.")
    left.add_paragraph("COMPÉTENCES")
    for count in (5, 5, 4, 4):
        left.add_paragraph("Compétences")
        for _ in range(count):
            paragraph = left.add_paragraph()
            paragraph.add_run("• ")
            paragraph.add_run("Python")
    left.add_paragraph("LANGUES")
    left.add_paragraph("English")
    left.add_paragraph("CENTRES D'INTÉRÊT")
    left.add_paragraph("Reading")
    left.add_paragraph("LIENS")
    left.add_paragraph("https://example.test")
    right.paragraphs[0].text = "PROJETS SÉLECTIONNÉS"
    for title in titles[:3]:
        paragraph = right.add_paragraph()
        paragraph.add_run(title)
        paragraph.add_run("  |  Python, tests")
        for text in ("Documented implementation.", "Documented evaluation."):
            paragraph = right.add_paragraph()
            paragraph.add_run("• ")
            paragraph.add_run(text)
    right.add_paragraph("FORMATION")
    right.add_paragraph("Documented education")
    document.save(template)
    generated = career / "CV/Generated"
    generated.mkdir()
    # Build a current variant through the production renderer; base template is deliberately legacy-shaped.
    from app.codex_workflows import CvContent
    from app.cv_documents import create_cv

    content = CvContent.model_validate({
        "headline": "Professional profile", "profile": "Documented candidate profile.",
        "skill_groups": [{"title": f"Skills {index}", "items": ["Python", "Testing", "Analysis"]} for index in range(4)],
        "projects": [{
            "project_id": title.lower().replace(" ", "-"), "title": title,
            "technologies": "Python", "description": "Documented project.",
            "bullets": ["Implemented a documented system.", "Evaluated its behavior."],
        } for title in titles[:3]],
    })
    create_cv(template, generated / "resume_A.docx", content)
    (career / "CV/CV_Variants.md").write_text(
        "# CV Variants\n\n## Variante A — Agentic AI\n\nProjets :\n\n"
        "1. Search Prototype\n2. Document Workflow\n3. Model Experiment\n",
        encoding="utf-8",
    )

    letter = career / "Cover Letter/Templates/lettre-motivation-template.docx"
    letter.parent.mkdir(parents=True)
    document = Document()
    for name in ("CANDIDATE_NAME", "CANDIDATE_ADDRESS", "CANDIDATE_ZIP_CODE", "CANDIDATE_EMAIL", "CANDIDATE_PHONE", "COMPANY_NAME", "COMPANY_ADDRESS", "COMPANY_ZIP_CODE"):
        document.add_paragraph("{{" + name + "}}")
    document.add_paragraph("Fait à {{CANDIDATE_CITY}}, le {{DATE}}")
    document.add_paragraph("Objet : Candidature - {{JOB_TITLE}}")
    for name in ("SALUTATION", "PARAGRAPH_1", "PARAGRAPH_2", "PARAGRAPH_3", "PARAGRAPH_4", "PARAGRAPH_5", "CLOSING", "SIGNATURE_NAME"):
        document.add_paragraph("{{" + name + "}}")
    document.save(letter)


def pytest_configure(config: pytest.Config) -> None:
    basetemp = config.getoption("basetemp")
    if basetemp and Path(basetemp).resolve().is_relative_to(REPO_ROOT):
        raise pytest.UsageError("pytest --basetemp doit être situé hors du dépôt")
    config._test_directory = TemporaryDirectory(prefix="job-tracker-tests-")
    root = Path(config._test_directory.name)
    repo = root / "knowledge-base"
    config._test_environment = pytest.MonkeyPatch()
    config._test_environment.setenv("KNOWLEDGE_BASE_PATH", str(repo))
    config._test_environment.setenv("DATABASE_PATH", str(root / "default.db"))
    config._test_environment.setenv("LOCALAPPDATA", str(root / "local-app-data"))
    config._test_environment.setenv("PYTHON_KEYRING_BACKEND", "keyring.backends.null.Keyring")
    _test_repository(repo)


def pytest_unconfigure(config: pytest.Config) -> None:
    if hasattr(config, "_test_environment"):
        config._test_environment.undo()
        config._test_directory.cleanup()


@pytest.fixture(scope="session")
def empty_database(pytestconfig):
    from app.main import init_db

    path = Path(pytestconfig._test_directory.name) / "empty.db"
    init_db(path)
    return path


@pytest.fixture(autouse=True)
def initialize_new_test_databases(monkeypatch, empty_database):
    from app import main

    initialize = main.init_db

    def initialize_database(path):
        if not path.exists():
            path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(empty_database, path)
        initialize(path)

    monkeypatch.setattr(main, "init_db", initialize_database)


@pytest.fixture(scope="session", autouse=True)
def no_repository_test_artifacts():
    before = _temporary_artifacts()
    yield
    created = _temporary_artifacts() - before
    assert not created, f"Artefacts temporaires laissés dans le dépôt: {sorted(map(str, created))}"
