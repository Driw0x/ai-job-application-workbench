from __future__ import annotations

import copy
import json
import os
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from pydantic import ValidationError

from app.codex_workflows import (
    CV_FINAL_REVIEW, CV_WRITING_RULES, CvContent, CvRefreshOutput, PreparationOutput,
    cv_refresh_prompt, preparation_context, preparation_prompt,
)
from app.cv_documents import (
    CvDocumentError, create_cv, cv_project_order, cv_variants, document_text, project_catalog,
    validate_cv_projects, validate_docx,
)
from app.main import add_event, connect
from scripts import refresh_cvs
from test_codex_integration import (
    FakeApiProvider, FakeKeyStore, FakeProviderError, assert_no_source_leak, client, cv_content, isolated_repo,
    preparation, shortlist, wait_for_preparation,
)


PROJECT_IDS = ["search-prototype", "document-workflow", "model-experiment", "planning-prototype", "image-classifier"]


@pytest.mark.parametrize("skill_count,bullet_count", [(3, 2), (4, 3), (5, 3)])
def test_cv_layout_rebuilds_sections_and_preserves_relevant_skills(tmp_path, skill_count, bullet_count):
    from docx import Document

    repo = isolated_repo(tmp_path)
    source = repo / "01_Career/CV/Templates/resume-template.docx"
    content = cv_content()
    for index, group in enumerate(content["skill_groups"]):
        group["items"] = [f"Compétence {index}-{item}" for item in range(skill_count)]
    for project in content["projects"]:
        project["bullets"] = ["Réalisation documentée."] * bullet_count
    content = CvContent.model_validate(content)
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    create_cv(source, first, content)
    create_cv(first, second, content)
    assert document_text(first) == document_text(second)
    document = Document(second)
    skills = document.tables[0].cell(0, 0).paragraphs
    texts = [p.text for p in skills]
    for group in content.skill_groups:
        start = texts.index(group.title) + 1
        retained = []
        while texts[start].startswith("• "):
            retained.append(texts[start][2:])
            start += 1
        assert 3 <= len(retained) <= skill_count
        assert retained == group.items[:len(retained)]
    if skill_count == 5:
        assert sum(text.startswith("• Compétence") for text in texts) < 20
    projects = document_text(second).split("PROJETS SÉLECTIONNÉS\n")[1].split("\nFORMATION")[0]
    assert projects.count("Projet technique documenté dans la Knowledge Base.") == len(content.projects)
    assert projects.count("• ") == len(content.projects) * bullet_count
    with ZipFile(source) as original, ZipFile(second) as generated:
        for name in original.namelist():
            if name != "word/document.xml":
                assert original.read(name) == generated.read(name)
    assert [(node.tag, dict(node.attrib)) for node in document.sections[0]._sectPr.iter()] == [
        (node.tag, dict(node.attrib)) for node in Document(source).sections[0]._sectPr.iter()
    ]
    original_texts = [p.text for p in Document(source).tables[0].cell(0, 0).paragraphs]
    assert texts[texts.index("LANGUES"):] == original_texts[original_texts.index("LANGUES"):]


def test_oversized_cv_keeps_existing_destination(tmp_path):
    repo = isolated_repo(tmp_path)
    source = repo / "01_Career/CV/Templates/resume-template.docx"
    content = cv_content()
    content["projects"][0]["description"] = "Description trop longue. " * 2000
    destination = tmp_path / "existing.docx"
    destination.write_bytes(b"original")
    with pytest.raises(CvDocumentError, match="Resume too long"):
        create_cv(source, destination, CvContent.model_validate(content), replace=True)
    assert destination.read_bytes() == b"original"


def test_old_cv_requires_adaptation_before_reuse(tmp_path):
    from app.cv_documents import validate_cv_layout

    repo = isolated_repo(tmp_path, with_variants=True)
    validate_cv_layout(repo / "01_Career/CV/Generated/resume_A.docx")
    with pytest.raises(CvDocumentError, match="choose ADAPT"):
        validate_cv_layout(repo / "01_Career/CV/Templates/resume-template.docx")


@pytest.fixture
def prepared_application(tmp_path, request):
    repo = isolated_repo(tmp_path, with_variants=True)
    selected = getattr(request, "param", 3)
    project_ids = PROJECT_IDS[:selected] if isinstance(selected, int) else selected
    payload = preparation() if project_ids == PROJECT_IDS[:3] else preparation("CREATE", project_ids=project_ids)
    sources = ["01_Career/Profil.md", *[
        f"02_Projects/{project['title']}/{project['title']}.md" for project in cv_content(project_ids)["projects"]
    ]]
    payload["local_sources"] = sources
    api, app = client(tmp_path, payload, repo_root=repo)
    app.state.applications_root = repo / "01_Career/Applications"
    application_id = shortlist(api)
    assert api.post(f"/applications/{application_id}/prepare-with-codex").status_code == 200
    application = wait_for_preparation(api, application_id, "completed")
    payload = {"content": cv_content(project_ids), "local_sources": sources}
    payload["content"]["projects"][0]["bullets"][0] = "Développé un outil pour explorer des sources documentées."
    app.state.fake_codex.response = payload
    return api, app, application_id, Path(application["cv_path"])


def application_row(app, application_id):
    with connect(app.state.db_path) as db:
        return dict(db.execute("SELECT * FROM applications WHERE id=?", (application_id,)).fetchone())


def set_fields(app, application_id, **values):
    with connect(app.state.db_path) as db:
        assignments = ", ".join(f"{name}=?" for name in values)
        db.execute(f"UPDATE applications SET {assignments} WHERE id=?", [*values.values(), application_id])


def events(app, application_id):
    with connect(app.state.db_path) as db:
        return [dict(row) for row in db.execute("SELECT * FROM events WHERE application_id=? ORDER BY id", (application_id,))]


@pytest.mark.parametrize("model", [PreparationOutput, CvRefreshOutput])
def test_cv_output_schemas_allow_one_to_four_projects(model):
    definitions = model.model_json_schema()["$defs"]
    fields = [definitions["CvContent"]["properties"]["projects"]]
    if model is PreparationOutput:
        fields.extend(definitions["CvDecision"]["properties"][name] for name in ("project_set", "project_order"))
    for field in fields:
        assert field["minItems"] == 1 and field["maxItems"] == 4
    for count in range(1, 5):
        payload = preparation("CREATE", project_ids=PROJECT_IDS[:count])
        if model is CvRefreshOutput:
            payload = {"content": payload["cv"]["content"], "local_sources": payload["local_sources"]}
        model.model_validate(payload)


@pytest.mark.parametrize("model", [PreparationOutput, CvRefreshOutput])
@pytest.mark.parametrize("count", [0, 5])
def test_cv_outputs_reject_project_counts_outside_one_to_four(model, count):
    payload = preparation("CREATE", project_ids=PROJECT_IDS[:count])
    if model is CvRefreshOutput:
        payload = {"content": payload["cv"]["content"], "local_sources": payload["local_sources"]}
    with pytest.raises(ValidationError):
        model.model_validate(payload)


@pytest.mark.parametrize("invalid", ["duplicate_set", "duplicate_order", "different_count", "different_content", "duplicate_reusable"])
def test_preparation_rejects_incoherent_project_decisions(invalid):
    payload = preparation("CREATE", project_ids=PROJECT_IDS[:3] if invalid.startswith("duplicate") else PROJECT_IDS[:4])
    decision = payload["cv"]
    if invalid == "duplicate_set":
        decision["project_set"] = [*PROJECT_IDS[:3], PROJECT_IDS[0]]
    elif invalid == "duplicate_order":
        decision["project_order"] = [*PROJECT_IDS[:3], PROJECT_IDS[0]]
    elif invalid == "different_count":
        decision["project_order"] = PROJECT_IDS[:3]
    elif invalid == "different_content":
        decision["content"]["projects"].reverse()
    else:
        decision["reusable_content"]["projects"].append(copy.deepcopy(decision["reusable_content"]["projects"][0]))
    with pytest.raises(ValidationError):
        PreparationOutput.model_validate(payload)


@pytest.mark.parametrize("project_ids", [[], PROJECT_IDS, [*PROJECT_IDS[:3], PROJECT_IDS[0]]])
@pytest.mark.parametrize("separator", ["\n", "\n\n"])
def test_variant_manifest_rejects_invalid_project_selection(tmp_path, project_ids, separator):
    repo = isolated_repo(tmp_path, with_variants=True)
    catalog = project_catalog(repo)
    projects = separator.join(f"{index}. {catalog[project_id]}" for index, project_id in enumerate(project_ids, 1))
    manifest = repo / "01_Career/CV/CV_Variants.md"
    manifest.write_text(f"# CV Variants\n\n## Variante A — Test\n\nProjets :\n\n{projects}\n", encoding="utf-8")
    assert cv_variants(repo) == {}


def test_variant_manifest_parses_four_projects_with_blank_lines(tmp_path):
    repo = isolated_repo(tmp_path, with_variants=True)
    catalog = project_catalog(repo)
    projects = "\n\n".join(f"{index}. {catalog[project_id]}" for index, project_id in enumerate(PROJECT_IDS[:4], 1))
    manifest = repo / "01_Career/CV/CV_Variants.md"
    manifest.write_text(f"# CV Variants\n\n## Variante A — Test\n\nProjets :\n\n{projects}\n", encoding="utf-8")
    assert cv_variants(repo)["A"]["project_order"] == PROJECT_IDS[:4]


def test_cv_prompts_share_writing_rules_and_fixed_refresh_projects():
    offer = {"company": "Example", "position": "Stage IA", "description": "Développer un prototype."}
    order = [project["project_id"] for project in cv_content()["projects"]]
    assert CV_WRITING_RULES in preparation_prompt(offer, "facts")
    assert preparation_prompt(offer, "facts").endswith(CV_FINAL_REVIEW)
    for instruction in (
        "ni un README", "Maximum une métrique", "ROC-AUC 0,78", "Ne fonde pas une puce",
        "N'invente aucune compétence", "3 puces complémentaires",
    ):
        assert instruction in CV_WRITING_RULES
    for scope in (offer, None):
        prompt = cv_refresh_prompt(scope, "facts", "current CV", order)
        assert CV_WRITING_RULES in prompt
        assert json.dumps(order, ensure_ascii=False) in prompt
        assert "set ET l'ordre" in prompt
        assert "description et 2 à 3 puces par projet" in prompt
        assert "local_sources" in prompt
        assert prompt.endswith(CV_FINAL_REVIEW)


def test_reuse_copies_three_projects_with_two_bullets_each(prepared_application):
    _, app, _, source = prepared_application
    variant = cv_variants(app.state.repo_root)["A"]["path"]
    assert source != variant
    assert source.read_bytes() == variant.read_bytes()
    lines = document_text(source).splitlines()
    section = lines[lines.index("PROJETS SÉLECTIONNÉS") + 1:lines.index("FORMATION")]
    assert sum(line.startswith("• ") for line in section) == 6


@pytest.mark.parametrize("action", ["REUSE", "ADAPT", "CREATE"])
def test_historical_preparation_uses_editorial_rules_for_every_cv_action(tmp_path, action):
    repo = isolated_repo(tmp_path, with_variants=True)
    api, app = client(tmp_path, preparation(action), repo_root=repo)
    application_id = shortlist(api)
    assert api.post(f"/applications/{application_id}/prepare-with-codex").status_code == 200
    application = wait_for_preparation(api, application_id, "completed")
    prompt = app.state.fake_codex.runs[-1][0][0]
    assert CV_WRITING_RULES in prompt
    assert prompt.endswith(CV_FINAL_REVIEW)
    event = next(event for event in events(app, application_id) if event["event_type"] == "CODEX_PREPARATION")
    assert json.loads(event["details"])["action"] == action
    validate_docx(Path(application["cv_path"]))


@pytest.mark.parametrize("prepared_application", [4], indirect=True)
def test_refresh_context_includes_selected_project_evidence_without_keyword_match(prepared_application):
    _, app, _, source = prepared_application
    order = cv_project_order(source, project_catalog(app.state.repo_root))
    context = preparation_context(
        app.state.repo_root, {"position": "Unrelated words", "description": ""},
        limit=0, selected_project_ids=order,
    )
    for project_id in order:
        title = project_catalog(app.state.repo_root)[project_id]
        assert f"SOURCE LOCALE: 02_Projects/{title}/{title}.md" in context






def test_project_validation_accepts_equivalent_dash_separated_titles(prepared_application, tmp_path):
    _, app, _, source = prepared_application
    content = cv_content()
    content["projects"][1]["title"] = "Document — Workflow"
    candidate = tmp_path / "normalized-project-titles.docx"
    create_cv(source, candidate, CvContent.model_validate(content))
    order = [project["project_id"] for project in content["projects"]]
    catalog = project_catalog(app.state.repo_root)
    validate_cv_projects(candidate, order, catalog)
    assert cv_project_order(candidate, catalog) == order


@pytest.mark.parametrize("prepared_application", [4], indirect=True)
def test_refresh_generator_uses_active_chatgpt_model_and_effort(prepared_application):
    api, app, application_id, source = prepared_application
    assert api.put("/codex/model", json={"model": "test-model", "effort": "high"}).status_code == 200
    order = cv_project_order(source, project_catalog(app.state.repo_root))
    output, metadata = app.state.generate_cv_refresh(application_row(app, application_id), source, order)
    assert isinstance(output, CvRefreshOutput)
    assert metadata == {"provider": "openai", "auth_mode": "chatgpt_oauth", "model": "test-model", "effort": "high"}
    arguments, options = app.state.fake_codex.runs[-1]
    assert arguments[1] == "test-model"
    assert arguments[2] == CvRefreshOutput.model_json_schema()
    assert options == {"effort": "high", "web_search": False}
    assert len(output.content.projects) == len(order) == 4


@pytest.mark.parametrize("provider_id", ["openai", "anthropic_api", "gemini_api", "deepseek_api"])
@pytest.mark.parametrize("prepared_application", [4], indirect=True)
def test_refresh_generator_uses_active_api_provider_without_chatgpt_fallback(prepared_application, provider_id):
    api, app, application_id, source = prepared_application
    provider = FakeApiProvider(provider_id, app.state.fake_codex.response, web_search=False)
    app.state.api_key_store = FakeKeyStore()
    app.state.api_key_store.set(provider_id, "test-only-api-key-1234567890")
    app.state.provider_factories[provider_id] = lambda _: provider
    assert api.put("/ai/active", json={
        "provider_id": provider_id, "auth_mode": "api_key", "model": f"{provider_id}-test",
        "effort": "medium", "confirm_api_billing": True,
    }).status_code == 200
    previous_calls = app.state.fake_codex.calls
    order = cv_project_order(source, project_catalog(app.state.repo_root))
    output, metadata = app.state.generate_cv_refresh(application_row(app, application_id), source, order)
    assert metadata == {"provider": provider_id, "auth_mode": "api_key", "model": f"{provider_id}-test", "effort": "medium"}
    assert provider.runs[-1][0][1] == f"{provider_id}-test"
    assert provider.runs[-1][1] == {"effort": "medium", "web_search": False}
    assert app.state.fake_codex.calls == previous_calls
    assert len(output.content.projects) == len(order) == 4


@pytest.mark.parametrize("use_fallback", [False, True])
def test_refresh_preserves_historical_stage_override_fallback_and_usage(prepared_application, use_fallback):
    api, app, application_id, source = prepared_application
    provider_id = "anthropic_api"
    provider = FakeApiProvider(provider_id, app.state.fake_codex.response, web_search=False)
    app.state.api_key_store = FakeKeyStore()
    app.state.api_key_store.set(provider_id, "test-only-api-key-1234567890")
    app.state.provider_factories[provider_id] = lambda _: provider
    assert api.put("/ai/pipeline", json={
        "overrides": {"application_preparation": {
            "provider": provider_id, "model": f"{provider_id}-test", "effort": "medium",
        }},
        "ai_fallback": "chatgpt_oauth" if use_fallback else None,
        "confirm_api_billing": True,
    }).status_code == 200
    if use_fallback:
        provider.error = FakeProviderError(503)
    before = application_row(app, application_id)
    order = cv_project_order(source, project_catalog(app.state.repo_root))
    output, metadata = app.state.generate_cv_refresh(before, source, order)
    assert isinstance(output, CvRefreshOutput)
    assert provider.runs[-1][0][1] == f"{provider_id}-test"
    assert metadata == {
        "provider": "openai" if use_fallback else provider_id,
        "auth_mode": "chatgpt_oauth" if use_fallback else "api_key",
        "model": "test-model" if use_fallback else f"{provider_id}-test",
        "effort": "medium",
    }
    with connect(app.state.db_path) as db:
        usage = db.execute("SELECT * FROM ai_usage ORDER BY id DESC LIMIT 1").fetchone()
    assert (usage["provider"], usage["model"], usage["category"], usage["operation"]) == (
        metadata["provider"], metadata["model"], "DOCUMENT", "application",
    )
    assert application_row(app, application_id) == before


@pytest.mark.parametrize("status,prepared_application", [
    (status, count) for status, count in zip(sorted(refresh_cvs.REFRESH_STATUSES), [1, 2, 3, 4, 3])
], indirect=["prepared_application"])
def test_cv_refresh_preserves_entire_application_row_and_other_artifacts(prepared_application, tmp_path, status):
    _, app, application_id, source = prepared_application
    set_fields(app, application_id, status=status, notes="Notes de suivi", next_action="Action conservée", next_action_at="2027-02-01", interview_notes="Notes conservées")
    before = application_row(app, application_id)
    before_events = events(app, application_id)
    target = refresh_cvs.targets(app)[0]
    protected = refresh_cvs.other_artifacts(target)
    source_hash = refresh_cvs.fingerprint(source)
    variant = next(variant["path"] for variant in cv_variants(app.state.repo_root).values()
                   if variant["project_order"] == target["project_order"])
    variant_hash = refresh_cvs.fingerprint(variant)
    directory = tmp_path / "reviewed"
    record = refresh_cvs.stage(app, target, directory)
    assert refresh_cvs.fingerprint(source) == source_hash
    assert application_row(app, application_id) == before
    assert events(app, application_id) == before_events
    assert refresh_cvs.apply(app, target, directory) == "refreshed"
    assert application_row(app, application_id) == before
    assert refresh_cvs.other_artifacts(target) == protected
    assert refresh_cvs.fingerprint(variant) == variant_hash
    assert refresh_cvs.fingerprint(source) == record["after_sha256"] != source_hash
    validate_docx(source)
    assert cv_project_order(source, project_catalog(app.state.repo_root)) == target["project_order"]
    text = document_text(source)
    assert_no_source_leak(text)
    projects = text.split("PROJETS SÉLECTIONNÉS\n", 1)[1].split("\nFORMATION", 1)[0]
    assert sum(line.startswith("• ") for line in projects.splitlines()) == 2 * len(target["project_order"])
    after_events = events(app, application_id)
    assert after_events[:-1] == before_events
    assert after_events[-1]["event_type"] == "CV_REFRESHED"
    assert refresh_cvs.already_refreshed(app, refresh_cvs.targets(app)[0])
    assert refresh_cvs.apply(app, target, directory) == "unchanged"
    assert events(app, application_id) == after_events


def test_refresh_excludes_every_non_validation_or_send_status(prepared_application):
    _, app, application_id, source = prepared_application
    source_hash = refresh_cvs.fingerprint(source)
    for status in ("DETECTED", "SENT", "ACKNOWLEDGED", "HR_INTERVIEW", "TECH_INTERVIEW", "FINAL_INTERVIEW", "OFFER", "REJECTED", "WITHDRAWN", "IGNORED", "NO_RESPONSE"):
        set_fields(app, application_id, status=status)
        assert refresh_cvs.targets(app) == [], status
    assert refresh_cvs.fingerprint(source) == source_hash


@pytest.mark.parametrize("reason", ["sent_at", "sent_event", "deleted"])
def test_refresh_excludes_prepared_application_with_recorded_send_or_deletion(prepared_application, reason):
    _, app, application_id, _ = prepared_application
    if reason == "sent_event":
        with connect(app.state.db_path) as db:
            add_event(db, application_id, "SENT", "Envoi historique")
    else:
        set_fields(app, application_id, **{"sent_at" if reason == "sent_at" else "deleted_at": "2026-10-07T12:00:00+00:00"})
    assert refresh_cvs.targets(app) == []


def test_refresh_rejects_cv_shared_with_archived_application(prepared_application):
    api, app, _, source = prepared_application
    other = api.post("/applications", json={"company": "Other", "position": "Stage", "url": "https://example.test/other"}).json()["id"]
    set_fields(app, other, status="SENT", cv_path=str(source), sent_at="2026-10-07T12:00:00+00:00")
    with pytest.raises(ValueError, match="shared"):
        refresh_cvs.targets(app)


@pytest.mark.parametrize("prepared_application", [1, 2, 3, 4], indirect=True)
def test_variant_refresh_keeps_application_copy_and_manifest_unchanged(prepared_application, tmp_path):
    _, app, application_id, source = prepared_application
    set_fields(app, application_id, status="SENT", sent_at="2026-10-07T12:00:00+00:00")
    order = cv_project_order(source, project_catalog(app.state.repo_root))
    target = next(target for target in refresh_cvs.targets(app, include_variants=True) if target["project_order"] == order)
    archived = source.read_bytes()
    before = application_row(app, application_id)
    before_events = events(app, application_id)
    manifest = app.state.repo_root / "01_Career/CV/CV_Variants.md"
    manifest_before = manifest.read_bytes()
    directory = tmp_path / "reviewed"
    refresh_cvs.stage(app, target, directory)
    assert refresh_cvs.apply(app, target, directory) == "refreshed"
    assert source.read_bytes() == archived
    assert application_row(app, application_id) == before
    assert events(app, application_id) == before_events
    assert manifest.read_bytes() == manifest_before
    assert cv_project_order(target["source"], project_catalog(app.state.repo_root)) == order
    assert cv_project_order(source, project_catalog(app.state.repo_root)) == order


def test_variant_refresh_rejects_global_file_referenced_by_archived_application(prepared_application):
    _, app, application_id, _ = prepared_application
    variant = app.state.repo_root / "01_Career/CV/Generated/resume_A.docx"
    set_fields(app, application_id, status="SENT", cv_path=str(variant), sent_at="2026-10-07T12:00:00+00:00")
    with pytest.raises(ValueError):
        refresh_cvs.targets(app, include_variants=True)


def test_refresh_rechecks_shared_paths_before_applying_reviewed_cv(prepared_application, tmp_path):
    api, app, _, source = prepared_application
    target = refresh_cvs.targets(app)[0]
    directory = tmp_path / "reviewed"
    refresh_cvs.stage(app, target, directory)
    original = source.read_bytes()
    other = api.post("/applications", json={"company": "Other", "position": "Stage", "url": "https://example.test/late-share"}).json()["id"]
    set_fields(app, other, status="SENT", cv_path=str(source), sent_at="2026-10-07T12:00:00+00:00")
    with pytest.raises(ValueError):
        refresh_cvs.apply(app, target, directory)
    assert source.read_bytes() == original


def test_refresh_rejects_preparation_started_after_staging(prepared_application, tmp_path):
    _, app, application_id, source = prepared_application
    set_fields(app, application_id, status="SHORTLISTED")
    target = refresh_cvs.targets(app)[0]
    directory = tmp_path / "reviewed"
    refresh_cvs.stage(app, target, directory)
    original = source.read_bytes()
    with connect(app.state.db_path) as db:
        add_event(db, application_id, "PREPARATION_STARTED", "Préparation concurrente")
    with pytest.raises(ValueError):
        refresh_cvs.apply(app, target, directory)
    assert source.read_bytes() == original


@pytest.mark.parametrize("event_type", ["PREPARATION_QUEUED", "PREPARATION_STARTED"])
def test_refresh_import_and_cli_preserve_active_preparation(prepared_application, monkeypatch, event_type):
    _, app, application_id, source = prepared_application
    set_fields(app, application_id, status="SHORTLISTED")
    with connect(app.state.db_path) as db:
        add_event(db, application_id, event_type, "Préparation active")
    before = application_row(app, application_id)
    before_events = events(app, application_id)
    original = source.read_bytes()
    subprocess.run(
        [sys.executable, "-c",
         "import sqlite3,sys; connect=sqlite3.connect; "
         "sqlite3.connect=lambda *args,**kwargs: connect(sys.argv[1],**kwargs); "
         "import app.main; import scripts.refresh_cvs", str(app.state.db_path)],
        env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"},
        cwd=Path(__file__).resolve().parents[1],
        capture_output=True, text=True, check=True,
    )
    assert application_row(app, application_id) == before
    assert events(app, application_id) == before_events
    monkeypatch.setattr(sys, "argv", [
        "refresh_cvs.py", "--database", str(app.state.db_path),
        "--knowledge-base", str(app.state.repo_root), "--applications-root", str(app.state.applications_root),
    ])
    with pytest.raises(ValueError, match="Preparation running"):
        refresh_cvs.main()
    assert application_row(app, application_id) == before
    assert events(app, application_id) == before_events
    assert source.read_bytes() == original


@pytest.mark.parametrize("change", ["source", "offer", "status", "notes"])
def test_refresh_rejects_changes_after_staging(prepared_application, tmp_path, change):
    _, app, application_id, source = prepared_application
    target = refresh_cvs.targets(app)[0]
    directory = tmp_path / "reviewed"
    refresh_cvs.stage(app, target, directory)
    if change == "source":
        source.write_bytes(source.read_bytes() + b"\nchanged")
    elif change == "offer":
        path = Path(application_row(app, application_id)["offer_path"])
        path.write_text(path.read_text(encoding="utf-8") + "\nNew offer facts", encoding="utf-8")
    else:
        set_fields(app, application_id, **{change: "SENT" if change == "status" else "New notes"})
    source_hash = refresh_cvs.fingerprint(source)
    before = application_row(app, application_id)
    with pytest.raises(ValueError, match="changed"):
        refresh_cvs.apply(app, target, directory)
    assert refresh_cvs.fingerprint(source) == source_hash
    assert application_row(app, application_id) == before
    assert not any(event["event_type"] == "CV_REFRESHED" for event in events(app, application_id))


@pytest.mark.parametrize("invalid", ["source", "order", "additional_project", "missing_project", "path", "structure"])
def test_invalid_refresh_output_keeps_original_cv_and_application(prepared_application, tmp_path, invalid):
    _, app, application_id, source = prepared_application
    payload = copy.deepcopy(app.state.fake_codex.response)
    if invalid == "source":
        payload["local_sources"] = ["01_Career/absent.md"]
    elif invalid == "order":
        payload["content"]["projects"].reverse()
    elif invalid == "additional_project":
        payload["content"]["projects"].append(cv_content([PROJECT_IDS[3]])["projects"][0])
    elif invalid == "missing_project":
        payload["content"]["projects"].pop()
    elif invalid == "path":
        payload["content"]["profile"] += " 02_Projects/private.md"
    else:
        payload["content"]["projects"][0]["bullets"].pop()
    app.state.fake_codex.response = payload
    target = refresh_cvs.targets(app)[0]
    source_hash = refresh_cvs.fingerprint(source)
    before = application_row(app, application_id)
    protected = refresh_cvs.other_artifacts(target)
    with pytest.raises(ValueError):
        refresh_cvs.stage(app, target, tmp_path / "reviewed")
    assert refresh_cvs.fingerprint(source) == source_hash
    assert application_row(app, application_id) == before
    assert refresh_cvs.other_artifacts(target) == protected


def test_refresh_restores_original_when_audit_write_fails(prepared_application, tmp_path, monkeypatch):
    _, app, application_id, source = prepared_application
    target = refresh_cvs.targets(app)[0]
    directory = tmp_path / "reviewed"
    refresh_cvs.stage(app, target, directory)
    original = source.read_bytes()
    before = application_row(app, application_id)
    def fail_event(*args, **kwargs):
        raise RuntimeError("Audit unavailable")
    monkeypatch.setattr(refresh_cvs, "add_event", fail_event)
    with pytest.raises(RuntimeError, match="Audit unavailable"):
        refresh_cvs.apply(app, target, directory)
    assert source.read_bytes() == original
    assert application_row(app, application_id) == before
    assert not any(event["event_type"] == "CV_REFRESHED" for event in events(app, application_id))


def test_refresh_expands_legacy_three_projects_with_grounded_fourth(prepared_application, tmp_path):
    _, app, application_id, source = prepared_application
    before = application_row(app, application_id)
    target = refresh_cvs.targets(app)[0]
    protected = refresh_cvs.other_artifacts(target)
    payload = app.state.fake_codex.response
    payload["content"] = cv_content(PROJECT_IDS[:4])
    payload["local_sources"].append("02_Projects/Planning Prototype/Planning Prototype.md")
    directory = tmp_path / "reviewed-four"
    record = refresh_cvs.stage(app, target, directory)
    assert record["project_order"] == PROJECT_IDS[:4]
    assert cv_project_order(source, project_catalog(app.state.repo_root)) == PROJECT_IDS[:3]
    assert refresh_cvs.apply(app, target, directory) == "refreshed"
    assert cv_project_order(source, project_catalog(app.state.repo_root)) == PROJECT_IDS[:4]
    assert application_row(app, application_id) == before
    assert refresh_cvs.other_artifacts(target) == protected


def test_english_resume_headings_support_four_grounded_projects(tmp_path):
    from docx import Document
    from app.cv_documents import CV_HEADINGS, validate_cv_layout

    repo = isolated_repo(tmp_path)
    source = repo / "01_Career/CV/Templates/resume-template.docx"
    document = Document(source)
    headings = {legacy: english for english, legacy in CV_HEADINGS.items()}
    for row in document.tables[0].rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if paragraph.text in headings:
                    paragraph.text = headings[paragraph.text]
    document.save(source)
    output = tmp_path / "english.docx"
    create_cv(source, output, CvContent.model_validate(cv_content(PROJECT_IDS[:4])))
    validate_cv_layout(output)
    assert "SELECTED PROJECTS" in document_text(output)
    assert cv_project_order(output, project_catalog(repo)) == PROJECT_IDS[:4]


def test_unique_user_named_resume_template_works_without_code_changes(tmp_path):
    from app.cv_documents import cv_template_path

    repo = isolated_repo(tmp_path)
    preferred = repo / "01_Career/CV/Templates/resume-template.docx"
    template = preferred.with_name("candidate-layout.docx")
    preferred.rename(template)
    assert cv_template_path(repo) == template
    context = preparation_context(repo, {"position": "Engineering", "description": ""})
    assert "CONTENU DOCX: 01_Career/CV/Templates/candidate-layout.docx" in context
    api, _ = client(tmp_path, preparation("CREATE", project_ids=PROJECT_IDS[:4]), repo_root=repo)
    application_id = shortlist(api)
    assert api.post(f"/applications/{application_id}/prepare-with-codex").status_code == 200
    result = wait_for_preparation(api, application_id, "completed")
    assert cv_project_order(Path(result["cv_path"]), project_catalog(repo)) == PROJECT_IDS[:4]
    preferred.write_bytes(template.read_bytes())
    assert cv_template_path(repo) == preferred
