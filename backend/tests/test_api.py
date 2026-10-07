from hashlib import sha256
import os
from pathlib import Path
from shutil import copy2
import subprocess
import sqlite3

from fastapi.testclient import TestClient
from docx import Document
import pytest
from app import main

from app.main import (
    CV_ROOT, change_status, cleanup_inactive_application_artifacts, connect, create_app,
    detect_application_channel, init_db, reset_to_prepared,
)
from app.cv_pdf import PdfExportError, PdfExportUnavailableError
from app.cover_letters import BODY_STYLE


def client(tmp_path: Path) -> TestClient:
    return TestClient(create_app(tmp_path / "test.db", tmp_path / "applications"))


def create(application_client: TestClient, url: str = "https://example.com/job/?tracking=1") -> dict:
    response = application_client.post("/applications", json={
        "company": "Example Company", "position": "Stage IA", "url": url,
        "source": "Test", "location": "Paris", "description": "Agents",
    })
    assert response.status_code == 201
    return response.json()


def complete_preparation(database: Path, applications_root: Path, application_id: int) -> None:
    with connect(database) as db:
        company, position = db.execute(
            "SELECT company, position FROM applications WHERE id=?", (application_id,)
        ).fetchone()
    folder = applications_root / f"example-company-stage-ia-{application_id}"
    folder.mkdir(parents=True)
    documents = {
        "offer_path": ("offer.md", "# Offre\n\nDescription factuelle."),
        "analysis_path": ("analysis.md", "# Analyse\n\n## Resume decision\n\nREUSE."),
        "company_path": ("company.md", "# Entreprise\n\nDescription."),
        "interview_prep_path": ("interview-prep.md", f"# Interview preparation — {company}\n\n" + "\n\n".join(
            f"## {section}\n\nContenu."
            for section in (
                "Position", "Company overview", "Role domain", "Main responsibilities",
                "Profile match", "Projects to highlight", "Technical topics to review",
                "Likely technical questions", "Likely HR / motivation questions",
                "Questions for the company", "Company-specific topics", "Actual interview",
            )
        )),
    }
    paths = {}
    for field, (name, content) in documents.items():
        path = folder / name
        path.write_text(content, encoding="utf-8")
        paths[field] = str(path)
    letter = folder / "lettre-motivation.docx"
    document = Document()
    document.styles.add_style(BODY_STYLE, 1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Alex Example"
    table.cell(0, 1).text = company
    document.add_paragraph("Fait à Paris, le 04/10/2026")
    document.add_paragraph(f"Objet : Candidature - {position}")
    document.add_paragraph("Madame, Monsieur,")
    for index in range(5):
        document.add_paragraph(f"Paragraphe métier {index + 1} pour {company} et le stage IA.", style=BODY_STYLE)
    document.add_paragraph("Cordialement,")
    document.add_paragraph("Alex Example")
    document.save(letter)
    paths["cover_letter_path"] = str(letter)
    cv = folder / "cv" / "CV.docx"
    cv.parent.mkdir()
    copy2(next(CV_ROOT.glob("*.docx")), cv)
    paths["cv_path"] = str(cv)
    with connect(database) as db:
        assignments = ", ".join(f"{field}=?" for field in paths)
        db.execute(
            f"UPDATE applications SET status='PREPARED', company_postal_address='10 rue du Test\n75001 Paris', {assignments} WHERE id=?",
            [*paths.values(), application_id],
        )
        db.commit()


def test_creation_retrieval_events_and_deduplication(tmp_path: Path) -> None:
    api = client(tmp_path)
    application = create(api)
    assert application["status"] == "DETECTED"
    assert api.get(f"/applications/{application['id']}").json()["company"] == "Example Company"
    events = api.get(f"/applications/{application['id']}/events").json()
    assert events[0]["event_type"] == "OFFER_DETECTED"
    duplicate = api.post("/applications", json={"company": "Example Company", "position": "Stage IA", "url": "https://example.com/job", "source": "Autre"})
    assert duplicate.status_code == 409


def test_application_patch_rejects_document_paths(tmp_path: Path) -> None:
    api = client(tmp_path)
    application = api.post("/applications", json={
        "company": "Example", "position": "Stage IA", "url": "https://example.test/job",
    }).json()

    response = api.patch(
        f"/applications/{application['id']}",
        json={"cv_path": "01_Career/CV/Generated/other.docx"},
    )

    assert response.status_code == 422


def test_manual_creation_stays_detected_without_preparation(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "Applications"
    interview_root = tmp_path / "Interview Preparation"
    codex_calls = []
    api = TestClient(create_app(
        database,
        applications_root,
        codex_factory=lambda token: codex_calls.append(token),
    ))

    response = api.post("/applications", json={
        "company": "TEST JOB TRACKER",
        "position": "Stage IA Test",
        "url": "https://example.com/job-tracker-test",
        "location": "Paris",
    })

    assert response.status_code == 201
    application = response.json()
    assert application["status"] == "DETECTED"
    assert application["url"] == "https://example.com/job-tracker-test"
    assert not any(application[field] for field in (
        "cv_path", "interview_prep_path", "offer_path", "analysis_path",
        "company_path",
    ))
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM applications").fetchone()[0] == 1
        assert db.execute("SELECT status FROM applications").fetchone()[0] == "DETECTED"
        assert [row[0] for row in db.execute("SELECT event_type FROM events")] == ["OFFER_DETECTED"]
    assert api.get("/stats").json()["phases"]["search"] == 1
    assert codex_calls == []
    assert not applications_root.exists()
    assert not interview_root.exists()

    duplicate = api.post("/applications", json={
        "company": "TEST JOB TRACKER",
        "position": "Stage IA Test",
        "url": "https://example.com/job-tracker-test/?tracking=duplicate",
    })
    assert duplicate.status_code == 409


def test_required_validation_and_submission_flow(tmp_path: Path) -> None:
    db_path = tmp_path / "test.db"
    api = TestClient(create_app(db_path, tmp_path / "applications"))
    application_id = create(api)["id"]
    assert api.post(f"/applications/{application_id}/submission/confirm").status_code == 409
    assert api.post(f"/applications/{application_id}/status", json={"status": "SHORTLISTED"}).status_code == 200
    assert api.post(f"/applications/{application_id}/prepare", json={}).status_code == 410
    complete_preparation(db_path, tmp_path / "applications", application_id)
    assert api.post(f"/applications/{application_id}/status", json={"status": "SENT"}).status_code == 409
    assert api.post(f"/applications/{application_id}/validate").status_code == 200
    sent = api.post(f"/applications/{application_id}/submission/confirm").json()
    assert sent["status"] == "SENT"
    assert sent["sent_at"]


def test_missing_artifact_blocks_validation_with_clear_error(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "applications"
    api = TestClient(create_app(database, applications_root))
    application_id = create(api)["id"]
    complete_preparation(database, applications_root, application_id)
    (applications_root / f"example-company-stage-ia-{application_id}" / "analysis.md").unlink()

    result = api.post(f"/applications/{application_id}/validate")

    assert result.status_code == 409
    assert result.json()["detail"]["message"] == "Incomplete preparation"
    assert result.json()["detail"]["missing"] == ["analysis.md"]


def test_markdown_letter_alone_does_not_complete_preparation(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "applications"
    api = TestClient(create_app(database, applications_root))
    application_id = create(api)["id"]
    complete_preparation(database, applications_root, application_id)
    folder = applications_root / f"example-company-stage-ia-{application_id}"
    (folder / "lettre-motivation.docx").unlink()
    legacy = folder / "lettre-motivation.md"
    legacy.write_text("# Ancienne lettre", encoding="utf-8")
    with connect(database) as db:
        db.execute("UPDATE applications SET cover_letter_path=? WHERE id=?", (str(legacy), application_id))
        db.commit()

    result = api.post(f"/applications/{application_id}/validate")

    assert result.status_code == 409
    assert "lettre-motivation.docx" in result.json()["detail"]["invalid"]


@pytest.mark.parametrize("notes", ["Visioconférence avec l'équipe technique", r"C:\Users\stage", r"\1", "un\\antislash"])
def test_interview_update_reuses_preparation_file(tmp_path: Path, notes: str) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "applications"
    api = TestClient(create_app(database, applications_root))
    application_id = create(api)["id"]
    complete_preparation(database, applications_root, application_id)
    path = applications_root / f"example-company-stage-ia-{application_id}" / "interview-prep.md"

    result = api.patch(f"/applications/{application_id}", json={
        "interview_at": "2026-10-12T14:30", "interview_notes": notes,
    })

    assert result.status_code == 200
    content = path.read_text(encoding="utf-8")
    assert "## Actual interview" in content and "### Scheduled interview" in content
    assert "2026-10-12T14:30" in content and f"- Notes : {notes}" in content


@pytest.mark.parametrize("failure", ["encoding", "permission"])
def test_unreadable_analysis_is_invalid_in_list_and_detail(tmp_path, monkeypatch, failure):
    database = tmp_path / "test.db"
    applications_root = tmp_path / "applications"
    api = TestClient(create_app(database, applications_root))
    application_id = create(api)["id"]
    complete_preparation(database, applications_root, application_id)
    path = applications_root / f"example-company-stage-ia-{application_id}" / "analysis.md"
    if failure == "encoding":
        path.write_bytes(b"\xff\xfe")
    else:
        read_text = Path.read_text
        def denied_read(self, *args, **kwargs):
            if self == path:
                raise PermissionError("Analyse illisible")
            return read_text(self, *args, **kwargs)
        monkeypatch.setattr(Path, "read_text", denied_read)
    for response in (api.get("/applications"), api.get(f"/applications/{application_id}")):
        assert response.status_code == 200
        application = response.json()[0] if isinstance(response.json(), list) else response.json()
        assert application["artifacts"] == {"complete": False, "missing": [], "invalid": ["analysis.md"]}


def test_cors_exposes_download_headers_only(tmp_path):
    api = TestClient(create_app(tmp_path / "test.db", tmp_path / "applications"))
    response = api.get("/applications", headers={"Origin": "http://127.0.0.1:5173"})
    assert response.status_code == 200
    assert response.headers["access-control-expose-headers"] == "Content-Disposition, X-PDF-Warning"


def test_example_research_artifacts_are_complete_and_validation_is_idempotent(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "applications"
    api = TestClient(create_app(database, applications_root))
    application = api.post("/applications", json={
        "company": "Example Research Organization",
        "position": "Research internship — document retrieval",
        "url": "https://example.com/example-research",
    }).json()
    complete_preparation(database, applications_root, application["id"])

    first = api.get(f"/applications/{application['id']}").json()["artifacts"]
    second = api.get(f"/applications/{application['id']}").json()["artifacts"]

    assert first == second == {"complete": True, "missing": [], "invalid": []}
    assert not (applications_root / f"example-company-stage-ia-{application['id']}" / "application.md").exists()


def test_tracking_follows_workflow_without_codex(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    api = TestClient(create_app(database, tmp_path / "applications"))
    application = create(api, "https://jobs.smartrecruiters.com/company/123")
    application_id = application["id"]
    assert application["application_channel"] == "SmartRecruiters"
    assert application["next_action"] == "Decide whether this job is relevant"
    assert application["sent_at"] is None
    assert application["interview_at"] is None

    shortlisted = api.post(f"/applications/{application_id}/status", json={"status": "SHORTLISTED"}).json()
    assert shortlisted["next_action"] == "Wait for application preparation"
    complete_preparation(database, tmp_path / "applications", application_id)
    prepared = api.get(f"/applications/{application_id}").json()
    assert prepared["next_action"] == "Review analysis, resume and cover letter, then approve"
    incompatible = api.patch(
        f"/applications/{application_id}", json={"next_action": "Relancer l'entreprise"},
    ).json()
    assert incompatible["next_action"] == "Review analysis, resume and cover letter, then approve"
    awaiting = api.post(f"/applications/{application_id}/validate").json()
    assert awaiting["next_action"] == "Submit on the company website, then confirm submission"
    assert awaiting["sent_at"] is None

    sent = api.post(f"/applications/{application_id}/submission/confirm").json()
    assert sent["next_action"] == "Follow up if no response"
    assert sent["next_action_at"][:10] > sent["sent_at"][:10]
    interview = api.post(f"/applications/{application_id}/status", json={"status": "HR_INTERVIEW"}).json()
    assert interview["next_action"] == "Prepare for interview"
    assert interview["interview_at"] is None
    rejected = api.post(f"/applications/{application_id}/status", json={"status": "REJECTED"}).json()
    assert rejected["next_action"] is None
    assert rejected["next_action_at"] is None


def test_closed_tracking_and_channel_detection() -> None:
    assert detect_application_channel("https://welcomekit.co/jobs/42") == "WelcomeKit"
    assert detect_application_channel("https://www.welcometothejungle.com/fr/companies/x/jobs/y") == "Welcome to the Jungle"
    assert detect_application_channel("https://careers.example.com/job/42") == "Career website"
    assert detect_application_channel("https://example.com/offer/42") is None


def test_reset_to_prepared_is_safe_idempotent_and_preserves_artifacts(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    api = TestClient(create_app(database, tmp_path / "applications"))
    application_id = create(api, "https://stages.example.com/job/42")["id"]
    preserved = {
        "cv_variant": "D", "cv_path": "cv.docx", "offer_path": "offer.md",
        "analysis_path": "analysis.md", "company_path": "company.md",
        "notes": "Conserver",
    }
    with connect(database) as db:
        assignments = ", ".join(f"{field} = ?" for field in preserved)
        db.execute(
            f"UPDATE applications SET status = 'SUBMITTING', {assignments} WHERE id = ?",
            [*preserved.values(), application_id],
        )
        db.commit()
        first = reset_to_prepared(db, application_id)
        second = reset_to_prepared(db, application_id)
        reset_events = db.execute(
            "SELECT COUNT(*) FROM events WHERE application_id = ? AND event_type = 'RESET_TO_PREPARED'",
            (application_id,),
        ).fetchone()[0]
    assert first == second
    assert first["status"] == "PREPARED"
    assert first["sent_at"] is None
    assert first["next_action"] == "Review analysis, resume and cover letter, then approve"
    assert first["next_action_at"] is None
    assert first["application_channel"] == "Career website"
    assert reset_events == 1
    for field, value in preserved.items():
        assert first[field] == value


def test_reset_to_prepared_refuses_recorded_send(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    api = TestClient(create_app(database, tmp_path / "applications"))
    application_id = create(api)["id"]
    with connect(database) as db:
        db.execute("UPDATE applications SET status = 'SUBMITTING' WHERE id = ?", (application_id,))
        db.execute(
            "INSERT INTO events(application_id, event_type, details, created_at) VALUES (?, 'SENT', '', '2026-10-04')",
            (application_id,),
        )
        db.commit()
        try:
            reset_to_prepared(db, application_id)
        except Exception as error:
            assert getattr(error, "status_code", None) == 409
        else:
            raise AssertionError("Le reset devait être refusé")


def test_stats(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    api = TestClient(create_app(database, tmp_path / "applications"))
    statuses = [
        "DETECTED", "IGNORED", "SHORTLISTED", "PREPARED", "AWAITING_VALIDATION",
        "SENT", "ACKNOWLEDGED", "REJECTED", "WITHDRAWN",
    ]
    for index, status in enumerate(statuses):
        application_id = create(api, f"https://example.com/job/{index}")["id"]
        with connect(database) as db:
            db.execute("UPDATE applications SET status=? WHERE id=?", (status, application_id))
            db.commit()
    with connect(database) as db:
        db.execute(
            "INSERT INTO discovery_runs(started_at, status, model, rejected_count, rejection_reasons) VALUES (?, ?, ?, ?, ?)",
            ("2026-10-04T12:00:00+00:00", "SUCCESS", "test", 3, '[{"reason":"CLOSED"},{"reason":"INVALID"},{"reason":"INVALID_SCHEMA"}]'),
        )
        db.commit()
    stats = api.get("/stats").json()
    assert stats["new"] == 1
    assert stats["phases"] == {
        "search": 1,
        "ignored": 1,
        "validation": 2,
        "send": 1,
        "tracking": 2,
        "rejected": 1,
    }

    ignored_id = api.get("/applications", params={"status": "IGNORED"}).json()[0]["id"]
    assert api.post(f"/applications/{ignored_id}/restore").status_code == 200
    assert api.get("/stats").json()["phases"] == {
        "search": 2, "ignored": 0, "validation": 2, "send": 1, "tracking": 2, "rejected": 1,
    }


def test_document_routes_allow_only_associated_documents(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    api = TestClient(create_app(database, tmp_path / "applications"))
    application_id = create(api)["id"]
    readme = tmp_path / "applications" / "associated.md"
    readme.parent.mkdir(parents=True, exist_ok=True)
    readme.write_text("# Associated document", encoding="utf-8")
    with connect(database) as db:
        db.execute(
            "UPDATE applications SET offer_path=?, analysis_path=?, company_path=?, interview_prep_path=?, cover_letter_path=? WHERE id=?",
            (str(readme), str(readme), str(readme), str(readme), str(readme), application_id),
        )
        db.commit()
    for kind in ("offer", "analysis", "company", "interview-prep", "cover-letter"):
        assert api.get(f"/applications/{application_id}/files/{kind}").status_code == 200
    assert api.get(f"/applications/{application_id}/files/secrets").status_code == 404
    assert api.get(f"/applications/{application_id + 1}/files/analysis").status_code == 404


def test_missing_document_route_names_expected_file(tmp_path: Path) -> None:
    api = client(tmp_path)
    application_id = create(api)["id"]

    result = api.get(f"/applications/{application_id}/files/analysis")

    assert result.status_code == 404
    assert result.json()["detail"] == "Missing document : analysis.md"


def test_cv_preview_endpoint_is_inline_docx_and_rejects_unassociated_paths(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    api = TestClient(create_app(database, tmp_path / "applications"))
    application_id = create(api)["id"]
    assert api.get(f"/applications/{application_id}/cv").status_code == 404
    cv = next(CV_ROOT.glob("*.docx"))
    with connect(database) as db:
        db.execute("UPDATE applications SET cv_path=? WHERE id=?", (str(cv), application_id))
        db.commit()
    response = api.get(f"/applications/{application_id}/cv")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert not response.headers.get("content-disposition", "").startswith("attachment")
    assert response.content == cv.read_bytes()
    with connect(database) as db:
        db.execute("UPDATE applications SET cv_path=? WHERE id=?", ("../../outside.docx", application_id))
        db.commit()
    assert api.get(f"/applications/{application_id}/cv").status_code == 400


def test_cv_docx_download_returns_canonical_file_without_pdf_conversion(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    repo = tmp_path / "repo"
    cv = repo / "applications" / "cv" / "resume_Example_Research.docx"
    cv.parent.mkdir(parents=True)
    cv.write_bytes(b"canonical DOCX")
    converted = []
    api = TestClient(create_app(
        database, repo / "applications", repo_root=repo,
        pdf_converter=lambda *_: converted.append(True),
    ))
    application_id = create(api)["id"]
    with connect(database) as db:
        db.execute("UPDATE applications SET cv_path=? WHERE id=?", (str(cv.relative_to(repo)), application_id))
        db.commit()

    response = api.get(f"/applications/{application_id}/cv.docx")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert response.headers["content-disposition"] == 'attachment; filename="resume_Example_Research.docx"'
    assert sha256(response.content).digest() == sha256(cv.read_bytes()).digest()
    assert converted == []

    missing_id = create(api, "https://example.com/other-job")["id"]
    missing = api.get(f"/applications/{missing_id}/cv.docx")
    assert missing.status_code == 404
    assert missing.json()["detail"] == "No resume attached"
    with connect(database) as db:
        db.execute("UPDATE applications SET cv_path=? WHERE id=?", ("../../outside.docx", missing_id))
        db.commit()
    assert api.get(f"/applications/{missing_id}/cv.docx").status_code == 400


def test_cv_pdf_export_uses_associated_docx_and_cleans_temporary_files(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    repo = tmp_path / "repo"
    cv = repo / "applications" / "cv" / "resume_Test.docx"
    cv.parent.mkdir(parents=True)
    cv.write_bytes(b"DOCX source unchanged")
    generated_directories = []

    def convert(source: Path, destination: Path) -> None:
        assert source == cv
        generated_directories.append(destination.parent)
        destination.write_bytes(b"%PDF-1.7\n/Type /Page\n%%EOF")

    api = TestClient(create_app(database, repo / "applications", repo_root=repo, pdf_converter=convert))
    application_id = create(api)["id"]
    with connect(database) as db:
        db.execute("UPDATE applications SET cv_path=? WHERE id=?", (str(cv.relative_to(repo)), application_id))
        db.commit()

    response = api.get(f"/applications/{application_id}/cv.pdf")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == 'attachment; filename="resume_Test.pdf"'
    assert response.headers["x-pdf-page-count"] == "1"
    assert response.content.startswith(b"%PDF-")
    assert cv.read_bytes() == b"DOCX source unchanged"
    assert generated_directories and not generated_directories[0].exists()


def test_cover_letter_docx_and_pdf_endpoints_use_canonical_docx(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "applications"
    converted = []

    def convert(source: Path, destination: Path) -> None:
        converted.append(source)
        destination.write_bytes(b"%PDF-1.7\n/Type /Page\n%%EOF")

    api = TestClient(create_app(database, applications_root, repo_root=tmp_path, pdf_converter=convert))
    application_id = create(api)["id"]
    complete_preparation(database, applications_root, application_id)
    letter = applications_root / f"example-company-stage-ia-{application_id}" / "lettre-motivation.docx"

    docx = api.get(f"/applications/{application_id}/cover-letter.docx")
    assert converted == []
    pdf = api.get(f"/applications/{application_id}/cover-letter.pdf")

    assert docx.status_code == 200
    assert docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert docx.headers["content-disposition"] == 'attachment; filename="lettre-motivation.docx"'
    assert sha256(docx.content).digest() == sha256(letter.read_bytes()).digest()
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")
    assert pdf.headers["content-disposition"] == 'attachment; filename="lettre-motivation.pdf"'
    assert converted == [letter]

    failing = TestClient(create_app(
        database, applications_root, repo_root=tmp_path,
        pdf_converter=lambda *_: (_ for _ in ()).throw(PdfExportError("Conversion Office impossible.")),
    ))
    failed_pdf = failing.get(f"/applications/{application_id}/cover-letter.pdf")
    assert failed_pdf.status_code == 500
    assert failed_pdf.json()["detail"] == "Conversion Office impossible."

    letter.unlink()
    assert api.get(f"/applications/{application_id}/cover-letter.docx").status_code == 404


def test_cv_pdf_export_rejects_missing_invalid_and_unassociated_paths(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    repo = tmp_path / "repo"
    repo.mkdir()
    api = TestClient(create_app(database, repo / "applications", repo_root=repo, pdf_converter=lambda *_: None))
    application_id = create(api)["id"]
    assert api.get(f"/applications/{application_id + 1}/cv.pdf").status_code == 404
    assert api.get(f"/applications/{application_id}/cv.pdf").status_code == 404

    with connect(database) as db:
        db.execute("UPDATE applications SET cv_path=? WHERE id=?", ("../../outside.docx", application_id))
        db.commit()
    assert api.get(f"/applications/{application_id}/cv.pdf").status_code == 400

    wrong_type = repo / "cv.pdf"
    wrong_type.write_bytes(b"%PDF-")
    with connect(database) as db:
        db.execute("UPDATE applications SET cv_path=? WHERE id=?", (str(wrong_type), application_id))
        db.commit()
    assert api.get(f"/applications/{application_id}/cv.pdf").status_code == 404


def test_cv_pdf_export_reports_unavailable_and_conversion_errors(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    repo = tmp_path / "repo"
    cv = repo / "CV.docx"
    repo.mkdir()
    cv.write_bytes(b"DOCX")

    def unavailable(*_args) -> None:
        raise PdfExportUnavailableError(
            "Export PDF indisponible : Microsoft Word et LibreOffice ne sont pas installés "
            "ou accessibles sur cette machine."
        )

    api = TestClient(create_app(database, repo / "applications", repo_root=repo, pdf_converter=unavailable))
    application_id = create(api)["id"]
    with connect(database) as db:
        db.execute("UPDATE applications SET cv_path=? WHERE id=?", ("CV.docx", application_id))
        db.commit()
    response = api.get(f"/applications/{application_id}/cv.pdf")
    assert response.status_code == 503
    assert response.json()["detail"] == (
        "Export PDF indisponible : Microsoft Word et LibreOffice ne sont pas installés "
        "ou accessibles sur cette machine."
    )

    failing = TestClient(create_app(database, repo / "applications", repo_root=repo, pdf_converter=lambda *_: (_ for _ in ()).throw(PdfExportError("Conversion Office impossible."))))
    response = failing.get(f"/applications/{application_id}/cv.pdf")
    assert response.status_code == 500
    assert response.json()["detail"] == "Conversion Office impossible."


def test_cv_pdf_export_rejects_empty_and_invalid_pdf(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    repo = tmp_path / "repo"
    cv = repo / "CV.docx"
    repo.mkdir()
    cv.write_bytes(b"DOCX")

    def empty(_source: Path, destination: Path) -> None:
        destination.write_bytes(b"")

    api = TestClient(create_app(database, repo / "applications", repo_root=repo, pdf_converter=empty))
    application_id = create(api)["id"]
    with connect(database) as db:
        db.execute("UPDATE applications SET cv_path='CV.docx' WHERE id=?", (application_id,))
        db.commit()
    assert api.get(f"/applications/{application_id}/cv.pdf").json()["detail"] == "PDF conversion failed : Generated PDF is empty."

    def invalid(_source: Path, destination: Path) -> None:
        destination.write_bytes(b"not a pdf")

    invalid_api = TestClient(create_app(database, repo / "applications", repo_root=repo, pdf_converter=invalid))
    response = invalid_api.get(f"/applications/{application_id}/cv.pdf")
    assert response.status_code == 500
    assert response.json()["detail"] == "PDF conversion failed : invalid PDF signature."


def test_ignored_offer_can_be_restored_without_codex_or_data_loss(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "Applications"
    codex_calls = []
    api = TestClient(create_app(
        database,
        applications_root,
        codex_factory=lambda token: codex_calls.append(token),
    ))
    created = api.post("/applications", json={
        "company": "TEST IGNORED",
        "position": "Stage test ignore restore",
        "url": "https://example.com/test-ignore-restore",
        "source": "Test manuel",
        "location": "Paris",
        "description": "Données conservées",
    }).json()

    assert api.post(f"/applications/{created['id']}/restore").status_code == 409
    ignored = api.post(
        f"/applications/{created['id']}/ignore",
        json={"details": "Ignored by user"},
    )
    assert ignored.status_code == 200
    assert ignored.json()["status"] == "IGNORED"
    discarded = api.get("/applications", params={"status": "IGNORED"}).json()
    assert [item["id"] for item in discarded] == [created["id"]]
    assert discarded[0]["ignored_at"]
    assert api.get("/stats").json()["phases"]["search"] == 0
    assert api.post("/applications", json={
        "company": "Duplicate", "position": "Duplicate", "url": created["url"],
    }).status_code == 409

    restored = api.post(f"/applications/{created['id']}/restore")
    assert restored.status_code == 200
    assert restored.json()["status"] == "DETECTED"
    for field in ("company", "position", "url", "source", "location", "description"):
        assert restored.json()[field] == created[field]
    assert api.get("/applications", params={"status": "IGNORED"}).json() == []
    assert api.get("/applications", params={"status": "DETECTED"}).json()[0]["id"] == created["id"]
    assert api.post(f"/applications/{created['id']}/restore").status_code == 409
    events = [event["event_type"] for event in reversed(api.get(f"/applications/{created['id']}/events").json())]
    assert events == ["OFFER_DETECTED", "IGNORED", "RESTORED"]
    assert codex_calls == []
    assert not applications_root.exists()


def test_withdrawn_and_ignored_cleanup_generated_documents_only(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "Applications"
    api = TestClient(create_app(database, applications_root, repo_root=tmp_path))
    template = tmp_path / "CV-template.docx"
    template.write_bytes(b"source template")

    for index, target in enumerate(("WITHDRAWN", "IGNORED")):
        application_id = create(api, f"https://example.com/cleanup/{index}")["id"]
        complete_preparation(database, applications_root, application_id)
        folder = applications_root / f"example-company-stage-ia-{application_id}"
        cv = folder / "cv" / "CV.docx"
        letter = folder / "lettre-motivation.docx"
        cv.with_suffix(".pdf").write_bytes(b"%PDF cv")
        cv.with_name("CV-copy.docx").write_bytes(b"generated copy")
        letter.with_suffix(".pdf").write_bytes(b"%PDF letter")
        if target == "WITHDRAWN":
            with connect(database) as db:
                db.execute("UPDATE applications SET status='SENT' WHERE id=?", (application_id,))
                db.commit()
            response = api.post(
                f"/applications/{application_id}/status", json={"status": target}
            )
        else:
            response = api.post(
                f"/applications/{application_id}/ignore", json={"details": "Test"}
            )

        assert response.status_code == 200
        assert response.json()["status"] == target
        assert response.json()["cv_path"] is None
        assert response.json()["cover_letter_path"] is None
        assert not cv.exists() and not cv.with_suffix(".pdf").exists()
        assert not cv.with_name("CV-copy.docx").exists()
        assert not letter.exists() and not letter.with_suffix(".pdf").exists()
        assert (folder / "analysis.md").is_file()
        assert (folder / "company.md").is_file()
        assert (folder / "offer.md").is_file()
        assert (folder / "interview-prep.md").is_file()
        assert template.read_bytes() == b"source template"
        assert api.get(f"/applications/{application_id}/cv.docx").status_code == 404
        assert api.get(f"/applications/{application_id}/cover-letter.docx").status_code == 404
        events = api.get(f"/applications/{application_id}/events").json()
        assert target in {event["event_type"] for event in events}
        assert "ARTIFACTS_CLEANED" in {event["event_type"] for event in events}

        with connect(database) as db:
            second = cleanup_inactive_application_artifacts(
                db, application_id, tmp_path, applications_root
            )
            assert db.execute(
                "SELECT COUNT(*) FROM applications WHERE id=?", (application_id,)
            ).fetchone()[0] == 1
        assert second["deleted"] == []


def test_cleanup_preserves_outside_and_shared_files(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "Applications"
    api = TestClient(create_app(database, applications_root, repo_root=tmp_path))
    first = create(api, "https://example.com/cleanup/shared-1")["id"]
    second = create(api, "https://example.com/cleanup/shared-2")["id"]
    outside = tmp_path / "manual.docx"
    outside.write_bytes(b"manual")
    shared_folder = applications_root / f"example-company-stage-ia-{first}"
    shared_folder.mkdir(parents=True)
    shared = shared_folder / "lettre-motivation.docx"
    shared.write_bytes(b"shared")
    with connect(database) as db:
        db.execute(
            "UPDATE applications SET status='IGNORED', cv_path=?, cover_letter_path=? WHERE id=?",
            (str(outside), str(shared), first),
        )
        db.execute(
            "UPDATE applications SET cover_letter_path=? WHERE id=?", (str(shared), second)
        )
        db.commit()
        report = cleanup_inactive_application_artifacts(
            db, first, tmp_path, applications_root
        )
        cleaned = db.execute(
            "SELECT cv_path, cover_letter_path FROM applications WHERE id=?", (first,)
        ).fetchone()

    assert outside.read_bytes() == b"manual"
    assert shared.read_bytes() == b"shared"
    assert cleaned["cv_path"] is None and cleaned["cover_letter_path"] is None
    assert {item["reason"] for item in report["kept"]} == {
        "path not allowed", "shared file",
    }


@pytest.mark.parametrize("dry_run", [True, False])
def test_cleanup_preserves_junction_targets(tmp_path: Path, dry_run: bool) -> None:
    database = tmp_path / "test.db"
    applications_root = tmp_path / "Applications"
    api = TestClient(create_app(database, applications_root, repo_root=tmp_path))
    application_id = create(api)["id"]
    folder = applications_root / f"example-company-stage-ia-{application_id}"
    folder.mkdir(parents=True)
    normal = folder / "lettre-motivation.docx"
    normal.write_bytes(b"generated")
    outside = tmp_path / "outside"
    outside.mkdir()
    external = outside / "manual.docx"
    external.write_bytes(b"keep")
    junction = folder / "cv"
    if os.name == "nt":
        result = subprocess.run(
            ["cmd", "/c", "mklink", "/J", str(junction), str(outside)],
            capture_output=True,
        )
        if result.returncode:
            pytest.skip("Windows junction creation is unavailable in this environment")
    else:
        junction.symlink_to(outside, target_is_directory=True)

    with connect(database) as db:
        db.execute("UPDATE applications SET status='IGNORED' WHERE id=?", (application_id,))
        db.commit()
        report = cleanup_inactive_application_artifacts(
            db, application_id, tmp_path, applications_root, dry_run=dry_run,
        )

    assert str(normal.resolve()) in report["deleted"]
    assert normal.exists() is dry_run
    assert external.read_bytes() == b"keep"
    assert not any(Path(path).is_relative_to(outside.resolve()) for path in report["deleted"])
    assert {"path": str(external.resolve()), "reason": "path not allowed"} in report["kept"]


def test_deleted_offer_lifecycle_preserves_data_and_files(tmp_path: Path) -> None:
    database = tmp_path / "test.db"
    linked_file = tmp_path / "existing-offer.md"
    linked_file.write_text("keep", encoding="utf-8")
    codex_calls = []
    api = TestClient(create_app(
        database,
        tmp_path / "Applications",
        codex_factory=lambda token: codex_calls.append(token),
    ))
    response = api.post("/applications", json={
        "company": "TEST DELETE",
        "position": "Stage test suppression",
        "url": "https://example.com/test-delete",
        "source": "Test manuel",
        "location": "Paris",
        "description": "Vérification de la corbeille",
    })
    assert response.status_code == 201
    application = response.json()
    application_id = application["id"]
    with connect(database) as db:
        db.execute(
            "UPDATE applications SET notes = ?, why_relevant = ?, offer_path = ? WHERE id = ?",
            ("Notes conservées", "Pertinence conservée", str(linked_file), application_id),
        )
        db.commit()
        before = dict(db.execute("SELECT * FROM applications WHERE id = ?", (application_id,)).fetchone())

    assert api.delete(f"/applications/{application_id}/permanent").status_code == 409
    deleted = api.delete(f"/applications/{application_id}")
    assert deleted.status_code == 200
    assert deleted.json()["deleted_at"]
    assert api.delete(f"/applications/{application_id}").status_code == 409
    assert api.get("/applications").json() == []
    assert api.get(f"/applications/{application_id}").status_code == 404
    assert api.get("/stats").json()["phases"]["search"] == 0
    assert [item["id"] for item in api.get("/applications/deleted").json()] == [application_id]
    assert api.post("/applications", json={
        "company": "Duplicate", "position": "Duplicate", "url": "https://example.com/test-delete",
    }).status_code == 409
    with connect(database) as db:
        stored = dict(db.execute("SELECT * FROM applications WHERE id = ?", (application_id,)).fetchone())
        assert stored["deleted_at"]
        for field, value in before.items():
            if field not in {"deleted_at", "updated_at"}:
                assert stored[field] == value

    restored = api.post(f"/applications/{application_id}/restore")
    assert restored.status_code == 200
    assert restored.json()["deleted_at"] is None
    assert restored.json()["status"] == "DETECTED"
    assert api.post(f"/applications/{application_id}/restore").status_code == 409
    assert [item["id"] for item in api.get("/applications").json()] == [application_id]
    assert api.get("/applications/deleted").json() == []

    assert api.delete(f"/applications/{application_id}").status_code == 200
    permanent = api.delete(f"/applications/{application_id}/permanent")
    assert permanent.status_code == 200
    assert permanent.json()["preserved_files"] == [str(linked_file)]
    assert linked_file.read_text(encoding="utf-8") == "keep"
    assert api.get("/applications/deleted").json() == []
    assert api.post(f"/applications/{application_id}/restore").status_code == 404
    assert api.delete(f"/applications/{application_id}/permanent").status_code == 404
    with connect(database) as db:
        assert db.execute("SELECT 1 FROM applications WHERE id = ?", (application_id,)).fetchone() is None
        assert db.execute("SELECT 1 FROM events WHERE application_id = ?", (application_id,)).fetchone() is None
    assert create(api, "https://example.com/test-delete")["status"] == "DETECTED"
    assert codex_calls == []


def test_deleted_at_migration_is_idempotent_and_preserves_existing_rows(tmp_path: Path) -> None:
    database = tmp_path / "existing.db"
    init_db(database)
    with connect(database) as db:
        application_id = create(TestClient(create_app(database, tmp_path / "Applications")))["id"]
        db.execute("ALTER TABLE applications DROP COLUMN deleted_at")
        db.commit()

    init_db(database)
    init_db(database)
    with connect(database) as db:
        columns = [row[1] for row in db.execute("PRAGMA table_info(applications)")]
        row = db.execute("SELECT company, deleted_at FROM applications WHERE id = ?", (application_id,)).fetchone()
    assert columns.count("deleted_at") == 1
    assert tuple(row) == ("Example Company", None)


def test_schema_migration_rolls_back_all_alters_on_error(tmp_path, monkeypatch):
    database = tmp_path / "existing.db"
    api = TestClient(create_app(database, tmp_path / "Applications"))
    application_id = create(api)["id"]
    with connect(database) as db:
        db.execute("ALTER TABLE applications DROP COLUMN contract_type")
        db.execute("ALTER TABLE applications DROP COLUMN published_at")
        db.commit()
    original_connect = main.connect
    class FailingConnection(sqlite3.Connection):
        def execute(self, sql, *args, **kwargs):
            if sql.startswith("ALTER TABLE applications ADD COLUMN published_at"):
                raise sqlite3.OperationalError("Migration interrompue")
            return super().execute(sql, *args, **kwargs)
    monkeypatch.setattr(main, "connect", lambda path: sqlite3.connect(path, factory=FailingConnection))
    with pytest.raises(sqlite3.OperationalError, match="Migration interrompue"):
        init_db(database)
    with original_connect(database) as db:
        columns = {row[1] for row in db.execute("PRAGMA table_info(applications)")}
        assert "contract_type" not in columns and "published_at" not in columns
        assert db.execute("SELECT company FROM applications WHERE id=?", (application_id,)).fetchone()[0] == "Example Company"
    monkeypatch.setattr(main, "connect", original_connect)
    init_db(database)
    init_db(database)
    with connect(database) as db:
        row = db.execute("SELECT company, contract_type, published_at FROM applications WHERE id=?", (application_id,)).fetchone()
        assert tuple(row) == ("Example Company", None, None)
