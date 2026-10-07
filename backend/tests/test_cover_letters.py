import re
import sqlite3
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from docx import Document

from app.cover_letters import (
    BODY_STYLE, COVER_LETTER_PLACEHOLDERS, CoverLetterError, CoverLetterTemplateError,
    candidate_contact, cover_letter_date, cover_letter_template_path, create_cover_letter_docx,
    document_full_text, is_subject, migrate_markdown_cover_letter,
    migrate_cover_letter_docx, read_cover_letter_docx,
    missing_optional_cover_letter_values,
    replace_in_runs, validate_cover_letter_docx,
    validate_cover_letter_template,
)
from app.main import REPO_ROOT
from scripts.migrate_cover_letters import migrate_all


SOURCE = """# Lettre de motivation

## Date

Fait à Example City, le 04/10/2026

## Objet

**Objet : Candidature — Stage IA**

Madame, Monsieur,

Je candidate au stage IA proposé par Example.

Mon parcours couvre intelligence artificielle et décision.

Mes projets démontrent préparation de données et évaluation.

Ces acquis correspondent aux missions du stage IA.

Je serais disponible pour un entretien.

Cordialement,

Alex Example
"""


def test_markdown_migration_is_valid_idempotent_and_deletes_only_after_validation(tmp_path: Path) -> None:
    markdown = tmp_path / "lettre-motivation.md"
    docx = tmp_path / "lettre-motivation.docx"
    markdown.write_text(SOURCE, encoding="utf-8")

    assert migrate_markdown_cover_letter(
        REPO_ROOT, markdown, docx, "Example", "Stage IA",
        "10 rue du Test\n75001 Paris",
    )
    assert markdown.exists()
    assert validate_cover_letter_docx(docx, "Example", "Alex Example") > 0
    original = docx.read_bytes()
    assert not migrate_markdown_cover_letter(
        REPO_ROOT, markdown, docx, "Example", "Stage IA",
        "10 rue du Test\n75001 Paris", remove_source=True,
    )
    assert docx.read_bytes() == original
    assert not markdown.exists()


def test_failed_markdown_migration_preserves_source(tmp_path: Path) -> None:
    markdown = tmp_path / "lettre-motivation.md"
    markdown.write_text("contenu non structuré", encoding="utf-8")

    with pytest.raises(CoverLetterError):
        migrate_markdown_cover_letter(
            REPO_ROOT, markdown, tmp_path / "lettre-motivation.docx",
            "Example", "Stage IA", "10 rue du Test\n75001 Paris", remove_source=True,
        )

    assert markdown.exists()
    assert not (tmp_path / "lettre-motivation.docx").exists()


def test_real_template_respects_contract() -> None:
    validate_cover_letter_template(cover_letter_template_path(REPO_ROOT))


def test_real_profile_provides_structured_candidate_contact() -> None:
    contact = candidate_contact(REPO_ROOT)

    assert all((
        contact.name,
        contact.postal_address,
        contact.postal_code,
        contact.city,
        contact.email,
        contact.phone,
    ))
    assert contact.location == "Example City"
    assert contact.city == "Example City"


def test_real_template_uses_candidate_contact_from_profile() -> None:
    contact = candidate_contact(REPO_ROOT)
    paragraphs = [f"Paragraphe {index} pour Example et le Stage IA." for index in range(1, 6)]

    with TemporaryDirectory() as directory:
        output = Path(directory) / "lettre-test.docx"
        create_cover_letter_docx(
            REPO_ROOT, output, "Example", "Stage IA", "10 rue du Test\n75001 Paris",
            paragraphs, created_on=date(2026, 10, 4),
        )
        text = document_full_text(Document(output))

    assert "{{CANDIDATE_" not in text
    assert "{{DATE}}" not in text
    assert not any(f"{{{{{name}}}}}" in text for name in COVER_LETTER_PLACEHOLDERS)
    assert all(value in text for value in (
        contact.name,
        contact.postal_address,
        contact.postal_code,
        contact.city,
        contact.email,
        contact.phone,
    ))
    assert "Fait à Example City, le 04/10/2026" in text
    assert re.findall(r"\b\d{2}/\d{2}/\d{4}\b", text) == ["04/10/2026"]
    assert text.index(contact.name) < text.index(contact.postal_address) < text.index(contact.postal_code)
    assert text.index(contact.postal_code) < text.index(contact.email) < text.index(contact.phone)
    assert missing_optional_cover_letter_values("10 rue du Test\n75001 Paris") == ()


@pytest.mark.parametrize(("postal_address", "expected"), (
    ("75001 Paris", ({"field": "COMPANY_ADDRESS", "label": "Company address"},)),
    ("10 rue du Test\nParis", ({"field": "COMPANY_ZIP_CODE", "label": "Company postal code"},)),
    (None, (
        {"field": "COMPANY_ADDRESS", "label": "Company address"},
        {"field": "COMPANY_ZIP_CODE", "label": "Company postal code"},
    )),
))
def test_optional_company_address_values_are_empty_without_residual_placeholders(
    tmp_path: Path, postal_address: str | None, expected: tuple[dict[str, str], ...],
) -> None:
    output = tmp_path / "lettre-test.docx"
    paragraphs = [f"Paragraphe {index} pour Example et le Stage IA." for index in range(1, 6)]

    create_cover_letter_docx(
        REPO_ROOT, output, "Example", "Stage IA", postal_address, paragraphs,
    )

    assert missing_optional_cover_letter_values(postal_address) == expected
    assert not re.search(r"\{\{[^{}]+\}\}", document_full_text(Document(output)))


def test_required_company_name_still_blocks_cover_letter_generation(tmp_path: Path) -> None:
    output = tmp_path / "lettre-test.docx"
    paragraphs = [f"Paragraphe {index} pour Example et le Stage IA." for index in range(1, 6)]

    with pytest.raises(CoverLetterError, match="COMPANY_NAME"):
        create_cover_letter_docx(
            REPO_ROOT, output, " ", "Stage IA", None, paragraphs,
        )

    assert not output.exists()


@pytest.mark.parametrize("company", ("{{COMPANY_NAME}}", "{{DATE}}", "Example {{UNKNOWN}}"))
def test_company_placeholders_are_rejected_without_changing_existing_letter(
    tmp_path: Path, company: str,
) -> None:
    output = tmp_path / "lettre-test.docx"
    original = b"existing letter"
    output.write_bytes(original)
    paragraphs = [f"Paragraphe {index} pour le Stage IA." for index in range(1, 6)]

    with pytest.raises(CoverLetterTemplateError, match="placeholders forbidden.*COMPANY_NAME"):
        create_cover_letter_docx(REPO_ROOT, output, company, "Stage IA", None, paragraphs)

    assert output.read_bytes() == original
    assert list(tmp_path.iterdir()) == [output]


def test_candidate_contact_reports_missing_required_profile_field(tmp_path: Path) -> None:
    profile = tmp_path / "01_Career" / "Profil.md"
    profile.parent.mkdir(parents=True)
    profile.write_text(
        """- **Nom** : Alex Example
- **Code postal** : 12345
- **Ville** : Example City
- **Email** : alex@example.test
- **Téléphone** : 07 00 00 00 00
- **Localisation** : Example City
""",
        encoding="utf-8",
    )

    with pytest.raises(CoverLetterError, match="Address"):
        candidate_contact(tmp_path)


def test_missing_template_is_explicit(tmp_path: Path) -> None:
    with pytest.raises(CoverLetterTemplateError, match="COVER_LETTER_TEMPLATE_INVALID"):
        validate_cover_letter_template(tmp_path / "missing.docx")


def test_template_reports_missing_and_duplicate_placeholders(tmp_path: Path) -> None:
    source = Document(cover_letter_template_path(REPO_ROOT))
    source.paragraphs[0].runs[0].text = "placeholder absent"
    source.add_paragraph("{{DATE}}")
    path = tmp_path / "invalid.docx"
    source.save(path)

    with pytest.raises(CoverLetterTemplateError) as caught:
        validate_cover_letter_template(path)

    assert caught.value.missing_placeholders == ("CANDIDATE_NAME",)
    assert caught.value.duplicate_placeholders == ("DATE",)
    assert "COVER_LETTER_TEMPLATE_INVALID" in str(caught.value)


def test_replacement_handles_placeholder_split_across_runs(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("{{CANDIDATE_")
    paragraph.add_run("NAME}}")
    path = tmp_path / "split-placeholder.docx"
    document.save(path)

    rendered = Document(path)
    replace_in_runs(rendered.paragraphs[0], "{{CANDIDATE_NAME}}", "Alex Example")

    assert rendered.paragraphs[0].text == "Alex Example"
    assert "{{CANDIDATE_NAME}}" not in document_full_text(rendered)


@pytest.mark.parametrize("split_runs", (False, True))
def test_replacement_replaces_each_occurrence_and_preserves_surrounding_text(split_runs: bool) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    parts = (
        ("Bonjour {{COMPANY_", "NAME}} et {{COMPANY_", "NAME}}.")
        if split_runs else ("Bonjour {{COMPANY_NAME}} et {{COMPANY_NAME}}.",)
    )
    for part in parts:
        paragraph.add_run(part)

    replace_in_runs(paragraph, "{{COMPANY_NAME}}", "Example")

    assert paragraph.text == "Bonjour Example et Example."


@pytest.mark.parametrize(("original", "value"), (
    ("{{COMPANY_NAME}}", "{{COMPANY_NAME}}"),
    ("{{COMPANY_NAME}}", "Example {{COMPANY_NAME}}"),
    ("{{COMPANY_NAME}}COMPANY_NAME}}", "{{"),
))
def test_replacement_rejects_values_that_do_not_remove_a_placeholder(original: str, value: str) -> None:
    document = Document()
    paragraph = document.add_paragraph(original)

    with pytest.raises(CoverLetterTemplateError, match="no progress"):
        replace_in_runs(paragraph, "{{COMPANY_NAME}}", value)

    assert paragraph.text == original


@pytest.mark.parametrize("value", ("Objet:", "Objet :", "Objet\N{NO-BREAK SPACE}:"))
def test_subject_accepts_optional_normal_or_non_breaking_space(value: str) -> None:
    assert is_subject(value)


def legacy_letter(path: Path, paragraph_count: int = 5) -> None:
    document = Document()
    document.styles.add_style(BODY_STYLE, 1)
    document.add_paragraph("Fait à Example City, le 04/10/2026")
    document.add_paragraph("Objet : Candidature — Stage IA")
    document.add_paragraph("Madame, Monsieur,")
    for index in range(paragraph_count):
        document.add_paragraph(
            f"Paragraphe historique {index + 1} pour Example et le Stage IA.",
            style=BODY_STYLE,
        )
    document.add_paragraph("Cordialement,")
    document.add_paragraph("Alex Example")
    document.save(path)


def test_docx_migration_preserves_text_date_and_is_idempotent(tmp_path: Path) -> None:
    path = tmp_path / "lettre-motivation.docx"
    legacy_letter(path)
    paragraphs = read_cover_letter_docx(path).paragraphs

    assert migrate_cover_letter_docx(
        REPO_ROOT, path, "Example", "Stage IA", "10 rue du Test\n75001 Paris"
    )
    first = path.read_bytes()
    assert read_cover_letter_docx(path).paragraphs == paragraphs
    assert "04/10/2026" in "\n".join(p.text for p in Document(path).paragraphs)
    assert not migrate_cover_letter_docx(
        REPO_ROOT, path, "Example", "Stage IA", "10 rue du Test\n75001 Paris"
    )
    assert path.read_bytes() == first


def test_ambiguous_docx_migration_preserves_original(tmp_path: Path) -> None:
    path = tmp_path / "lettre-motivation.docx"
    legacy_letter(path, paragraph_count=4)
    original = path.read_bytes()

    with pytest.raises(CoverLetterError, match="Invalid cover-letter DOCX structure"):
        migrate_cover_letter_docx(
            REPO_ROOT, path, "Example", "Stage IA", "10 rue du Test\n75001 Paris"
        )

    assert path.read_bytes() == original


def test_batch_migration_reports_invalid_date_and_continues(tmp_path: Path) -> None:
    invalid = tmp_path / "invalid-date.docx"
    valid = tmp_path / "valid-date.docx"
    legacy_letter(invalid)
    legacy_letter(valid)
    document = Document(invalid)
    document.paragraphs[0].text = "Fait à Example City, le 31/02/2026"
    document.save(invalid)
    original = invalid.read_bytes()
    paragraphs = read_cover_letter_docx(valid).paragraphs

    with pytest.raises(CoverLetterError, match="Invalid historical date.*31/02/2026"):
        cover_letter_date(invalid)

    database = tmp_path / "migration.db"
    with sqlite3.connect(database) as connection:
        connection.execute(
            "CREATE TABLE applications (id INTEGER PRIMARY KEY, company TEXT, position TEXT, "
            "company_postal_address TEXT, cover_letter_path TEXT)"
        )
        connection.executemany(
            "INSERT INTO applications VALUES (?, ?, ?, ?, ?)",
            [
                (1, "Example", "Stage IA", "10 rue du Test\n75001 Paris", str(invalid)),
                (2, "Example", "Stage IA", "10 rue du Test\n75001 Paris", str(valid)),
            ],
        )
    connection.close()

    report = migrate_all(repo=REPO_ROOT, database_path=database)

    assert report["applications_total"] == report["letters_found"] == 2
    assert report["letters_failed"] == 1
    assert report["letters_migrated"] == report["letters_validated"] == 1
    assert report["failures"][0]["application_id"] == 1
    assert "Invalid historical date" in report["failures"][0]["reason"]
    assert invalid.read_bytes() == original
    assert read_cover_letter_docx(valid).paragraphs == paragraphs
    assert cover_letter_date(valid) == date(2026, 10, 4)

    migrated = valid.read_bytes()
    repeated = migrate_all(repo=REPO_ROOT, database_path=database)
    assert repeated["applications_total"] == repeated["letters_found"] == 2
    assert repeated["letters_migrated"] == 0
    assert repeated["letters_unchanged"] == repeated["letters_validated"] == repeated["letters_failed"] == 1
    assert repeated["failures"] == report["failures"]
    assert valid.read_bytes() == migrated
    assert invalid.read_bytes() == original


def test_final_validation_rejects_corruption_and_remaining_placeholder(tmp_path: Path) -> None:
    corrupt = tmp_path / "corrupt.docx"
    corrupt.write_bytes(b"not a docx")
    with pytest.raises(CoverLetterError, match="Unreadable"):
        validate_cover_letter_docx(corrupt)

    path = tmp_path / "placeholder.docx"
    legacy_letter(path)
    document = Document(path)
    document.paragraphs[3].add_run(" {{UNRESOLVED}}")
    document.save(path)
    with pytest.raises(CoverLetterError, match="unresolved placeholder"):
        validate_cover_letter_docx(path)


@pytest.mark.parametrize("address,street,postal", [
    ("12 Example Street\nLondon SW1A 1AA", "12 Example Street", "London SW1A 1AA"),
    ("123 Example Street\nExample City CA 94105", "123 Example Street", "Example City CA 94105"),
    ("12 Example Street\nExample City K1A 0B1", "12 Example Street", "Example City K1A 0B1"),
    ("12 Example Street\n101 Example City", "12 Example Street", "101 Example City"),
])
def test_international_company_postal_blocks(address, street, postal):
    from app.company_profiles import valid_postal_address
    from app.cover_letters import split_postal_address

    assert valid_postal_address(address)
    assert split_postal_address(address) == (street, postal)
    assert not valid_postal_address("123 Example Street\nExample City")
