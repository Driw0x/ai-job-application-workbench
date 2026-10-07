from __future__ import annotations

import hashlib
import os
import re
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph

from .cv_documents import CvDocumentError, document_text
from .company_profiles import POSTAL_CODE_PATTERN


BODY_STYLE = "Cover Letter Body"  # Legacy letters only; templates own new document styles.
COVER_LETTER_TEMPLATE = Path("01_Career/Cover Letter/Templates/lettre-motivation-template.docx")
COVER_LETTER_PLACEHOLDERS = (
    "CANDIDATE_NAME",
    "CANDIDATE_ADDRESS",
    "CANDIDATE_ZIP_CODE",
    "CANDIDATE_EMAIL",
    "CANDIDATE_PHONE",
    "COMPANY_NAME",
    "COMPANY_ADDRESS",
    "COMPANY_ZIP_CODE",
    "CANDIDATE_CITY",
    "DATE",
    "JOB_TITLE",
    "SALUTATION",
    "PARAGRAPH_1",
    "PARAGRAPH_2",
    "PARAGRAPH_3",
    "PARAGRAPH_4",
    "PARAGRAPH_5",
    "CLOSING",
    "SIGNATURE_NAME",
)
OPTIONAL_COVER_LETTER_VALUES = (
    "COMPANY_ADDRESS",
    "COMPANY_ZIP_CODE",
)
MISSING_INFORMATION_LABELS = {
    "COMPANY_ADDRESS": "Company address",
    "COMPANY_ZIP_CODE": "Company postal code",
}
PLACEHOLDER = re.compile(r"\{\{[^{}]+\}\}")
DATE_PATTERN = re.compile(r"\b\d{2}/\d{2}/\d{4}\b")
TEMPLATE_HASH_PREFIX = "cover-letter-template-sha256:"
INTERNAL_SOURCE = re.compile(
    r"(?:Source(?:s)? locale(?:s)?|SOURCE LOCALE)|(?:01_Career|02_Projects|03_Knowledge)[\\/]",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class CandidateContact:
    name: str
    location: str | None
    postal_address: str
    postal_code: str
    city: str
    email: str
    phone: str


@dataclass(frozen=True)
class CoverLetterContent:
    subject: str
    salutation: str
    paragraphs: tuple[str, ...]
    closing: str
    signature_name: str


class CoverLetterError(ValueError):
    pass


class CoverLetterTemplateError(CoverLetterError):
    code = "COVER_LETTER_TEMPLATE_INVALID"

    def __init__(
        self,
        message: str,
        *,
        missing_placeholders: Iterable[str] = (),
        duplicate_placeholders: Iterable[str] = (),
    ) -> None:
        self.missing_placeholders = tuple(missing_placeholders)
        self.duplicate_placeholders = tuple(duplicate_placeholders)
        details = []
        if self.missing_placeholders:
            details.append(f"missing_placeholders={list(self.missing_placeholders)!r}")
        if self.duplicate_placeholders:
            details.append(f"duplicate_placeholders={list(self.duplicate_placeholders)!r}")
        suffix = f" ({', '.join(details)})" if details else ""
        super().__init__(f"{self.code}: {message}{suffix}")


def cover_letter_template_path(repo: Path) -> Path:
    preferred = repo / "01_Career/Cover Letter/Templates/cover-letter-template.docx"
    return preferred if preferred.is_file() else repo / COVER_LETTER_TEMPLATE


def candidate_contact(repo: Path) -> CandidateContact:
    """Load candidate contact details from the canonical Career profile."""
    career = repo / "01_Career"
    profile_path = career / "Profile.md" if (career / "Profile.md").is_file() else career / "Profil.md"
    profile = profile_path.read_text(encoding="utf-8")
    location = _match(profile, r"\*\*(?:Location|Localisation)\*\*\s*:\s*([^\r\n]+)", 1)
    fields = {
        "Nom": _match(profile, r"\*\*(?:Name|Nom)\*\*\s*:\s*([^\r\n]+)", 1),
        "Adresse": _match(profile, r"\*\*(?:Address|Adresse(?: postale)?)\*\*\s*:\s*([^\r\n]+)", 1),
        "Code postal": _match(profile, r"\*\*(?:Postal code|Code postal)\*\*\s*:\s*([^\r\n]+)", 1),
        "Ville": _match(profile, r"\*\*(?:City|Ville)\*\*\s*:\s*([^\r\n]+)", 1),
        "Email": _match(profile, r"\*\*Email\*\*\s*:\s*([^\r\n]+)", 1),
        "Téléphone": _match(profile, r"\*\*(?:Phone|Téléphone)\*\*\s*:\s*([^\r\n]+)", 1),
    }
    missing = [
        {
            "Nom": "Name", "Adresse": "Address", "Code postal": "Postal code",
            "Ville": "City", "Téléphone": "Phone",
        }.get(name, name)
        for name, value in fields.items() if not value
    ]
    if missing:
        raise CoverLetterError(
            f"Missing candidate contact details in the profile : {', '.join(missing)}"
        )
    return CandidateContact(
        name=fields["Nom"] or "",
        location=location,
        postal_address=fields["Adresse"] or "",
        postal_code=fields["Code postal"] or "",
        city=fields["Ville"] or "",
        email=fields["Email"] or "",
        phone=fields["Téléphone"] or "",
    )


def default_cover_letter_content(
    position: str, paragraphs: list[str], language: str, signature_name: str
) -> CoverLetterContent:
    salutation, closing, label, prefix = (
        ("Dear Hiring Team,", "Sincerely,", "Subject", "Application")
        if language == "en"
        else ("Madame, Monsieur,", "Cordialement,", "Objet", "Candidature")
    )
    return CoverLetterContent(
        subject=f"{label} : {prefix} - {normalize_position(position)}",
        salutation=salutation,
        paragraphs=tuple(paragraph.strip() for paragraph in paragraphs),
        closing=closing,
        signature_name=signature_name,
    )


def create_cover_letter_docx(
    repo: Path,
    destination: Path,
    company: str,
    position: str,
    postal_address: str | None,
    paragraphs: list[str],
    language: str = "fr",
    created_on: date | None = None,
    salutation: str | None = None,
    closing: str | None = None,
) -> int:
    contact = candidate_contact(repo)
    content = default_cover_letter_content(position, paragraphs, language, contact.name)
    content = CoverLetterContent(
        subject=content.subject,
        salutation=(salutation or content.salutation).strip(),
        paragraphs=content.paragraphs,
        closing=(closing or content.closing).strip(),
        signature_name=contact.name,
    )
    validate_cover_letter_content(content, company, position)
    company_address, company_zip_code = split_postal_address(postal_address)
    values = {
        "CANDIDATE_NAME": contact.name,
        "CANDIDATE_ADDRESS": contact.postal_address,
        "CANDIDATE_ZIP_CODE": contact.postal_code,
        "CANDIDATE_EMAIL": contact.email,
        "CANDIDATE_PHONE": contact.phone,
        "COMPANY_NAME": company.strip(),
        "COMPANY_ADDRESS": company_address,
        "COMPANY_ZIP_CODE": company_zip_code,
        "CANDIDATE_CITY": contact.city,
        "DATE": (created_on or date.today()).strftime("%d/%m/%Y"),
        "JOB_TITLE": normalize_position(position),
        "SALUTATION": content.salutation,
        **{f"PARAGRAPH_{index}": value for index, value in enumerate(content.paragraphs, 1)},
        "CLOSING": content.closing,
        "SIGNATURE_NAME": content.signature_name,
    }
    required_values = (
        "CANDIDATE_NAME", "CANDIDATE_ADDRESS", "CANDIDATE_ZIP_CODE",
        "CANDIDATE_EMAIL", "CANDIDATE_PHONE", "CANDIDATE_CITY", "COMPANY_NAME",
        "DATE", "JOB_TITLE", "SALUTATION", "CLOSING", "SIGNATURE_NAME",
        "PARAGRAPH_1", "PARAGRAPH_2", "PARAGRAPH_3", "PARAGRAPH_4", "PARAGRAPH_5",
    )
    missing = [name for name in required_values if not values[name].strip()]
    if missing:
        raise CoverLetterError(f"Missing cover-letter values : {', '.join(missing)}")
    return render_cover_letter_from_template(
        cover_letter_template_path(repo), destination, values, company, position, contact.name
    )


def missing_optional_cover_letter_values(postal_address: str | None) -> tuple[dict[str, str], ...]:
    company_address, company_zip_code = split_postal_address(postal_address)
    values = {
        "COMPANY_ADDRESS": company_address,
        "COMPANY_ZIP_CODE": company_zip_code,
    }
    return tuple(
        {"field": name, "label": MISSING_INFORMATION_LABELS[name]}
        for name in OPTIONAL_COVER_LETTER_VALUES
        if not values[name].strip()
    )


def validate_cover_letter_template(template_path: Path) -> None:
    if not template_path.is_file():
        raise CoverLetterTemplateError(f"template not found : {template_path}")
    try:
        document = Document(template_path)
    except Exception as error:
        raise CoverLetterTemplateError(f"template Unreadable DOCX : {error}") from error
    text = "\n".join(paragraph.text for paragraph in iter_document_paragraphs(document))
    counts = {name: text.count(f"{{{{{name}}}}}") for name in COVER_LETTER_PLACEHOLDERS}
    missing = [name for name, count in counts.items() if count == 0]
    duplicates = [name for name, count in counts.items() if count > 1]
    if missing or duplicates:
        raise CoverLetterTemplateError(
            "placeholder contract violated",
            missing_placeholders=missing,
            duplicate_placeholders=duplicates,
        )


def render_cover_letter_from_template(
    template_path: Path,
    output_path: Path,
    values: dict[str, str],
    company: str | None = None,
    position: str | None = None,
    signature_name: str | None = None,
) -> int:
    validate_cover_letter_template(template_path)
    missing_values = [name for name in COVER_LETTER_PLACEHOLDERS if name not in values]
    if missing_values:
        raise CoverLetterTemplateError(
            "missing replacement values", missing_placeholders=missing_values
        )
    placeholder_values = [
        name for name in COVER_LETTER_PLACEHOLDERS if PLACEHOLDER.search(str(values[name]))
    ]
    if placeholder_values:
        raise CoverLetterTemplateError(
            f"placeholders forbidden in values : {placeholder_values}"
        )
    document = Document(template_path)
    for paragraph in iter_document_paragraphs(document):
        for name in COVER_LETTER_PLACEHOLDERS:
            replace_in_runs(paragraph, f"{{{{{name}}}}}", str(values[name]))
    remaining = sorted(set(PLACEHOLDER.findall(document_full_text(document))))
    if remaining:
        raise CoverLetterTemplateError(f"unresolved placeholders : {remaining}")
    document.core_properties.comments = TEMPLATE_HASH_PREFIX + template_sha256(template_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix=".docx", delete=False) as handle:
            temporary = Path(handle.name)
        document.save(temporary)
        count = validate_cover_letter_docx(temporary, company, signature_name, position)
        temporary.replace(output_path)
        return count
    except (OSError, CoverLetterError) as error:
        if temporary:
            temporary.unlink(missing_ok=True)
        if isinstance(error, CoverLetterError):
            raise
        raise CoverLetterError(f"Unable to write cover letter DOCX : {error}") from error


def validate_cover_letter_content(
    content: CoverLetterContent, company: str, position: str
) -> int:
    if len(content.paragraphs) != 5 or any(not value.strip() for value in content.paragraphs):
        raise CoverLetterError("Cover letter must contain exactly five body paragraphs")
    text = "\n".join((content.subject, content.salutation, *content.paragraphs, content.closing, content.signature_name))
    words = re.findall(r"\b[\wÀ-ÿ'-]+\b", "\n".join(content.paragraphs))
    if len(words) > 420:
        raise CoverLetterError("Cover letter exceeds 420 words")
    if "à compléter" in text.casefold() or PLACEHOLDER.search(text):
        raise CoverLetterError("Cover letter contains a forbidden placeholder")
    if INTERNAL_SOURCE.search(text):
        raise CoverLetterError("Knowledge Base reference forbidden in cover letter")
    keywords = [word.casefold() for word in re.findall(r"[\wÀ-ÿ-]{6,}", position)]
    if keywords and not any(word in text.casefold() for word in keywords):
        raise CoverLetterError("Cover letter does not match the position")
    return len(words)


def validate_cover_letter_docx(
    path: Path,
    company: str | None = None,
    signature_name: str | None = None,
    position: str | None = None,
) -> int:
    if not path.is_file() or path.suffix.lower() != ".docx" or path.stat().st_size == 0:
        raise CoverLetterError("Cover letter is not a valid DOCX file")
    try:
        document = Document(path)
    except Exception as error:
        raise CoverLetterError(f"Unreadable cover-letter DOCX : {error}") from error
    text = document_full_text(document)
    if PLACEHOLDER.search(text):
        raise CoverLetterError("Cover letter DOCX contains an unresolved placeholder")
    content = read_cover_letter_docx(path)
    if company and company.casefold() not in text.casefold():
        raise CoverLetterError("Recipient missing from cover-letter DOCX")
    if signature_name and signature_name.casefold() not in text.casefold():
        raise CoverLetterError("Signature missing from cover-letter DOCX")
    if position:
        normalized = normalize_position(position)
        if normalized.casefold() not in content.subject.casefold():
            raise CoverLetterError("Job title missing from cover-letter DOCX")
    if INTERNAL_SOURCE.search(text):
        raise CoverLetterError("Knowledge Base reference forbidden in cover letter DOCX")
    try:
        document_text(path)
    except CvDocumentError as error:
        raise CoverLetterError(str(error)) from error
    return len(re.findall(r"\b[\wÀ-ÿ'-]+\b", "\n".join(content.paragraphs)))


def read_cover_letter_docx(path: Path) -> CoverLetterContent:
    try:
        document = Document(path)
    except Exception as error:
        raise CoverLetterError(f"Unreadable cover-letter DOCX : {error}") from error
    paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    legacy_body = tuple(
        paragraph.text.strip() for paragraph in paragraphs if paragraph.style.name == BODY_STYLE
    )
    if len(legacy_body) == 5:
        first = next(index for index, paragraph in enumerate(paragraphs) if paragraph.style.name == BODY_STYLE)
        last = max(index for index, paragraph in enumerate(paragraphs) if paragraph.style.name == BODY_STYLE)
        if first < 1 or last + 2 >= len(paragraphs):
            raise CoverLetterError("Invalid cover-letter DOCX structure")
        subject = next(
            (paragraph.text.strip() for paragraph in paragraphs if is_subject(paragraph.text)), ""
        )
        return CoverLetterContent(
            subject=subject,
            salutation=paragraphs[first - 1].text.strip(),
            paragraphs=legacy_body,
            closing=paragraphs[last + 1].text.strip(),
            signature_name=paragraphs[last + 2].text.strip(),
        )
    subject_index = next((index for index, paragraph in enumerate(paragraphs) if is_subject(paragraph.text)), -1)
    tail = paragraphs[subject_index:] if subject_index >= 0 else []
    if len(tail) != 9:
        raise CoverLetterError("Invalid cover-letter DOCX structure")
    return CoverLetterContent(
        subject=tail[0].text.strip(),
        salutation=tail[1].text.strip(),
        paragraphs=tuple(paragraph.text.strip() for paragraph in tail[2:7]),
        closing=tail[7].text.strip(),
        signature_name=tail[8].text.strip(),
    )


def cover_letter_date(path: Path) -> date | None:
    document = Document(path)
    match = DATE_PATTERN.search(document_full_text(document))
    if not match:
        return None
    try:
        return datetime.strptime(match.group(), "%d/%m/%Y").date()
    except ValueError as error:
        raise CoverLetterError(f"Invalid historical date : {match.group()}") from error


def uses_current_cover_letter_template(path: Path, template_path: Path) -> bool:
    try:
        return Document(path).core_properties.comments == TEMPLATE_HASH_PREFIX + template_sha256(template_path)
    except Exception:
        return False


def migrate_cover_letter_docx(
    repo: Path,
    path: Path,
    company: str,
    position: str,
    postal_address: str | None,
) -> bool:
    """Reapply current template while preserving the five existing paragraphs."""
    template_path = cover_letter_template_path(repo)
    contact = candidate_contact(repo)
    if uses_current_cover_letter_template(path, template_path):
        validate_cover_letter_docx(path, company, contact.name, position)
        return False
    content = read_cover_letter_docx(path)
    existing_date = cover_letter_date(path)
    if existing_date is None:
        raise CoverLetterError("Ambiguous extraction: historical date not found")
    language = "en" if content.salutation.casefold().startswith("dear") else "fr"
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
            temporary = Path(handle.name)
        create_cover_letter_docx(
            repo, temporary, company, position, postal_address, list(content.paragraphs), language,
            existing_date, content.salutation, content.closing,
        )
        migrated = read_cover_letter_docx(temporary)
        if migrated.paragraphs != content.paragraphs:
            raise CoverLetterError("Historical text of five paragraphs was not preserved")
        if cover_letter_date(temporary) != existing_date:
            raise CoverLetterError("Historical date was not preserved")
        validate_cover_letter_docx(temporary, company, contact.name, position)
        os.replace(temporary, path)
        temporary = None
        return True
    finally:
        if temporary:
            temporary.unlink(missing_ok=True)


def parse_markdown_cover_letter(markdown: str) -> tuple[CoverLetterContent, date | None]:
    blocks = [block.strip() for block in re.split(r"\n\s*\n", markdown.strip()) if block.strip()]
    greeting_pattern = re.compile(r"^(?:Madame, Monsieur,|Dear Hiring Team,)$", re.IGNORECASE)
    closing_pattern = re.compile(r"^(?:Cordialement,|Sincerely,|Best regards,)$", re.IGNORECASE)
    greeting_index = next((index for index, block in enumerate(blocks) if greeting_pattern.match(block)), -1)
    if greeting_index < 0:
        raise CoverLetterError("Salutation not found in cover-letter Markdown")
    closing_index = next(
        (index for index in range(greeting_index + 1, len(blocks)) if closing_pattern.match(blocks[index])),
        -1,
    )
    if closing_index < 0 or closing_index + 1 >= len(blocks):
        raise CoverLetterError("Closing or signature not found in cover-letter Markdown")
    paragraphs = tuple(blocks[greeting_index + 1:closing_index])
    if len(paragraphs) != 5:
        raise CoverLetterError("Cover letter Markdown must contain five body paragraphs")
    subject_match = re.search(r"(?ms)^## Objet\s*\n+\*\*(.+?)\*\*", markdown)
    if not subject_match:
        raise CoverLetterError("Subject not found in cover-letter Markdown")
    date_match = re.search(r"(?ms)^## Date\s*\n+.*?(\d{2}/\d{2}/\d{4})", markdown)
    parsed_date = datetime.strptime(date_match.group(1), "%d/%m/%Y").date() if date_match else None
    return CoverLetterContent(
        subject=subject_match.group(1).strip(),
        salutation=blocks[greeting_index],
        paragraphs=paragraphs,
        closing=blocks[closing_index],
        signature_name=blocks[closing_index + 1].strip(),
    ), parsed_date


def migrate_markdown_cover_letter(
    repo: Path,
    markdown_path: Path,
    docx_path: Path,
    company: str,
    position: str,
    postal_address: str | None,
    remove_source: bool = False,
) -> bool:
    contact = candidate_contact(repo)
    if docx_path.exists():
        validate_cover_letter_docx(docx_path, company, contact.name)
        if remove_source:
            markdown_path.unlink(missing_ok=True)
        return False
    content, created_on = parse_markdown_cover_letter(markdown_path.read_text(encoding="utf-8"))
    language = "en" if content.salutation.casefold().startswith("dear") else "fr"
    create_cover_letter_docx(
        repo, docx_path, company, position, postal_address, list(content.paragraphs), language,
        created_on, content.salutation, content.closing,
    )
    validate_cover_letter_docx(docx_path, company, contact.name, position)
    if remove_source:
        markdown_path.unlink()
    return True


def split_postal_address(value: str | None) -> tuple[str, str]:
    lines = [line.strip() for line in (value or "").splitlines() if line.strip()]
    postal_index = next((
        index for index, line in reversed(list(enumerate(lines)))
        if (index > 0 or re.search(r"\b\d{5}\b", line))
        and re.search(POSTAL_CODE_PATTERN, line, re.IGNORECASE)
    ), -1)
    if postal_index < 0:
        return ("\n".join(lines), "")
    return ("\n".join(lines[:postal_index]), "\n".join(lines[postal_index:]))


def normalize_position(value: str) -> str:
    return re.sub(r"\s*(?:[-–]\s*)?\(?(?:F/H/N|H/F|F/H)\)?\s*$", "", value, flags=re.IGNORECASE).strip()


def iter_document_paragraphs(document: DocumentObject) -> Iterable[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        yield from iter_table_paragraphs(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from iter_table_paragraphs(table)


def iter_table_paragraphs(table: Table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def replace_in_runs(paragraph: Paragraph, token: str, value: str) -> None:
    while token in "".join(run.text for run in paragraph.runs):
        runs = paragraph.runs
        combined = "".join(run.text for run in runs)
        start = combined.index(token)
        end = start + len(token)
        updated = combined[:start] + value + combined[end:]
        if updated.count(token) >= combined.count(token):
            raise CoverLetterTemplateError("Placeholder replacement made no progress")
        offsets = []
        cursor = 0
        for run in runs:
            offsets.append((cursor, cursor + len(run.text)))
            cursor += len(run.text)
        start_run = next(index for index, (_, stop) in enumerate(offsets) if stop > start)
        end_run = next(index for index, (_, stop) in enumerate(offsets) if stop >= end)
        start_offset = start - offsets[start_run][0]
        end_offset = end - offsets[end_run][0]
        prefix = runs[start_run].text[:start_offset]
        suffix = runs[end_run].text[end_offset:]
        if start_run == end_run:
            runs[start_run].text = prefix + value + suffix
            continue
        runs[start_run].text = prefix + value
        for index in range(start_run + 1, end_run):
            runs[index].text = ""
        runs[end_run].text = suffix


def document_full_text(document: DocumentObject) -> str:
    return "\n".join(paragraph.text for paragraph in iter_document_paragraphs(document))


def template_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def is_subject(value: str) -> bool:
    return bool(re.match(r"^(?:Objet|Subject)\s*:", value.strip()))


def _match(text: str, pattern: str, group: int = 0) -> str | None:
    match = re.search(pattern, text, re.IGNORECASE)
    return match.group(group).strip() if match else None
