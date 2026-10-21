from __future__ import annotations

import re
import tempfile
import unicodedata
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
CV_HEADINGS = {
    "PROFILE": "PROFIL", "SKILLS": "COMPÉTENCES", "LANGUAGES": "LANGUES",
    "SELECTED PROJECTS": "PROJETS", "PROJETS SÉLECTIONNÉS": "PROJETS", "PROJECTS": "PROJETS",
    "EDUCATION": "FORMATION", "EXPERIENCE": "EXPÉRIENCES", "INTERESTS": "CENTRES D'INTÉRÊT", "LINKS": "LIENS",
}



class CvDocumentError(ValueError):
    pass


def canonical_project_id(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.casefold()).strip("-")


def normalized_project_text(value: str) -> str:
    return canonical_project_id(value).replace("-", "")


def project_catalog(repo: Path) -> dict[str, str]:
    root = repo / "02_Projects"
    return {
        canonical_project_id(path.name): path.name
        for path in sorted(root.iterdir())
        if path.is_dir() and any(path.glob("*.md"))
    } if root.is_dir() else {}


def cv_template_path(repo: Path) -> Path:
    directory = repo / "01_Career" / "CV" / "Templates"
    preferred = directory / "resume-template.docx"
    if preferred.is_file():
        return preferred
    candidates = sorted(directory.glob("*.docx"))
    return candidates[0] if len(candidates) == 1 else preferred


def cv_variants(repo: Path) -> dict[str, dict[str, object]]:
    manifest = repo / "01_Career" / "CV" / "CV_Variants.md"
    generated = repo / "01_Career" / "CV" / "Generated"
    if not manifest.is_file():
        return {}
    sections = re.split(r"(?m)^## (?:Variant|Variante) ", manifest.read_text(encoding="utf-8"))[1:]
    variants = {}
    for section in sections:
        lines = section.splitlines()
        variant = lines[0].split(" ", 1)[0]
        match = re.search(
            r"(?m)^(?:Projects|Projets) :\s*\n((?:\s*\d+\. [^\r\n]+(?:\n|$))+)",
            section,
        )
        path = generated / f"resume_{variant}.docx"
        if not path.is_file():
            candidates = sorted(generated.glob(f"*_{variant}.docx"))
            if len(candidates) == 1:
                path = candidates[0]
        if match and path.is_file():
            projects = [canonical_project_id(line.split(". ", 1)[1].strip()) for line in match[1].splitlines() if line.strip()]
            if 1 <= len(projects) <= 4 and len(set(projects)) == len(projects):
                variants[variant] = {"path": path, "project_order": projects, "project_set": frozenset(projects)}
    return variants


def next_variant_name(variants: dict[str, object]) -> str:
    used = set(variants)
    for code in range(ord("A"), ord("Z") + 1):
        if (candidate := chr(code)) not in used:
            return candidate
    raise CvDocumentError("No variant identifier available")


def register_variant(repo: Path, variant: str, project_order: list[str], catalog: dict[str, str], headline: str) -> None:
    manifest = repo / "01_Career" / "CV" / "CV_Variants.md"
    projects = "\n".join(f"{index}. {catalog[project_id]}" for index, project_id in enumerate(project_order, 1))
    entry = (
        f"\n## Variant {variant} — {headline}\n\nProjects :\n\n{projects}\n\n"
        "Purpose:\n\n- Reusable variant created by AI Job Application Workbench.\n"
    )
    current = manifest.read_text(encoding="utf-8") if manifest.is_file() else "# CV Variants\n"
    manifest.parent.mkdir(parents=True, exist_ok=True)
    manifest.write_text(current.rstrip() + "\n" + entry, encoding="utf-8")


def document_text(path: Path) -> str:
    try:
        with ZipFile(path) as archive:
            if archive.testzip() is not None or "word/document.xml" not in archive.namelist():
                raise CvDocumentError("Invalid or incomplete DOCX")
            root = ET.fromstring(archive.read("word/document.xml"))
    except (BadZipFile, KeyError, ET.ParseError, OSError) as error:
        raise CvDocumentError(f"Unreadable DOCX : {error}") from error
    return "\n".join(
        text for paragraph in root.iter(W + "p")
        if (text := "".join(node.text or "" for node in paragraph.iter(W + "t")).strip())
    )


def validate_docx(path: Path) -> None:
    if not path.is_file() or path.suffix.lower() != ".docx":
        raise CvDocumentError("Generated resume is not a DOCX file")
    if not document_text(path).strip():
        raise CvDocumentError("Generated resume is empty")


ATS_SECTIONS = ("PROFIL", "COMPÉTENCES", "PROJETS", "FORMATION", "EXPÉRIENCES", "LANGUES", "CENTRES D'INTÉRÊT")


def validate_cv_layout(path: Path) -> None:
    document = Document(path)
    lines = document_text(path).splitlines()
    try:
        positions = [lines.index(title) for title in ATS_SECTIONS if title in lines]
        if not all(title in lines for title in ("PROFIL", "COMPÉTENCES", "PROJETS", "FORMATION", "LANGUES")):
            raise ValueError
        if positions != sorted(positions) or document.tables:
            raise ValueError
        if any(document._element.iter(W + "txbxContent")):
            raise ValueError
        for section in document.sections:
            cols = section._sectPr.find(W + "cols")
            if cols is not None and cols.get(W + "num", "1") != "1":
                raise ValueError
            if abs(section.page_width.cm - 21) > .02 or abs(section.page_height.cm - 29.7) > .02:
                raise ValueError
    except ValueError as error:
        raise CvDocumentError("Legacy or incompatible resume layout: choose ADAPT for ATS layout") from error


def validate_cv_projects(
    path: Path, expected: list[str], catalog: dict[str, str], rendered_titles: dict[str, str] | None = None,
) -> None:
    lines = [CV_HEADINGS.get(line, line) for line in document_text(path).splitlines()]
    normalized = [canonical_project_id(line) for line in lines]
    try:
        start = next(index for index, title in enumerate(normalized) if title in {"projets", "projets-selectionnes"}) + 1
        end = normalized.index("formation", start)
    except (ValueError, StopIteration) as error:
        raise CvDocumentError("Resume projects section not found") from error
    section = normalized_project_text("\n".join(lines[start:end]))
    titles = {**catalog, **(rendered_titles or {})}
    found = [project_id for project_id, title in titles.items() if normalized_project_text(title) in section]
    if not 1 <= len(expected) <= 4 or len(found) != len(expected) or set(found) != set(expected):
        raise CvDocumentError("Resume must contain exactly the expected projects (1 to 4 distinct projects)")


def cv_project_order(path: Path, catalog: dict[str, str]) -> list[str]:
    lines = [CV_HEADINGS.get(line, line) for line in document_text(path).splitlines()]
    normalized = [canonical_project_id(line) for line in lines]
    try:
        start = next(index for index, title in enumerate(normalized) if title in {"projets", "projets-selectionnes"}) + 1
        end = normalized.index("formation", start)
    except (ValueError, StopIteration) as error:
        raise CvDocumentError("Resume projects section not found") from error
    order = []
    for line in lines[start:end]:
        normalized_line = normalized_project_text(line)
        matches = [key for key, title in catalog.items() if normalized_line.startswith(normalized_project_text(title))]
        if matches:
            order.append(max(matches, key=lambda key: len(catalog[key])))
    if not 1 <= len(order) <= 4 or len(set(order)) != len(order):
        raise CvDocumentError("Resume must contain exactly the expected projects (1 to 4 distinct projects)")
    validate_cv_projects(path, order, catalog)
    return order


def _cv_sections(source: Path) -> tuple[list[str], dict[str, list[str]]]:
    """Read legacy table cells and ATS paragraphs in their existing section order."""
    header, sections = [], {}
    current = header
    lines = document_text(source).splitlines()
    # Our visual model separates technologies; restore the semantic title for rendering.
    paragraphs = Document(source).paragraphs
    if any(p.style.name == "CV Project Tech" for p in paragraphs):
        lines = []
        for p in paragraphs:
            if not p.text.strip():
                continue
            if p.style.name == "CV Project Tech" and lines:
                lines[-1] += "  |  " + p.text.strip()
            else:
                lines.append(p.text.strip())
    for text in lines:
        heading = CV_HEADINGS.get(text, text)
        if heading in (*ATS_SECTIONS, "LIENS"):
            if heading in sections:
                raise CvDocumentError(f"Duplicate resume section: {heading}")
            current = sections.setdefault(heading, [])
        else:
            current.append(text)
    for required in ("PROFIL", "COMPÉTENCES", "PROJETS", "FORMATION", "LANGUES"):
        if required not in sections:
            raise CvDocumentError(f"Section '{required}' missing from resume template")
    return header, sections


def _compact_skills(lines: list[str]) -> list[str]:
    compact = []
    for text in lines:
        if text.startswith("• ") and compact:
            compact[-1] += (" : " if " : " not in compact[-1] else " · ") + text[2:]
        else:
            compact.append(text)
    return compact


def document_hyperlinks(path: Path) -> dict[str, str]:
    """Read visible links, including referenced headers and footers, excluding stale parts."""
    from docx.oxml.ns import qn

    document = Document(path)
    parts = [document.part]
    for section in document.sections:
        for reference in section._sectPr:
            if reference.tag in {W + "headerReference", W + "footerReference"}:
                parts.append(document.part.related_parts[reference.get(qn("r:id"))])
    result = {}
    for part in parts:
        for node in part.element.iter(W + "hyperlink"):
            relation = part.rels.get(node.get(qn("r:id")))
            label = "".join(t.text or "" for t in node.iter(W + "t"))
            if relation is not None and relation.is_external and label:
                if label in result and result[label] != relation.target_ref:
                    raise CvDocumentError(f"Ambiguous hyperlink label: {label}")
                result[label] = relation.target_ref
    return result


def convert_cv_to_ats(source: Path, destination: Path, content: object | None = None, replace: bool = False) -> None:
    """One renderer for CREATE, ADAPT, REUSE and content-preserving migration."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    if destination.exists() and not replace:
        raise CvDocumentError(f"Target resume already exists : {destination.name}")
    validate_docx(source)
    header, sections = _cv_sections(source)
    if len(header) < 3:
        raise CvDocumentError("Resume identity header is incomplete")
    if content is not None:
        header[2] = content.headline
        sections["PROFIL"] = [content.profile]
        sections["COMPÉTENCES"] = [group.title + " : " + " · ".join(group.items) for group in content.skill_groups]
        sections["PROJETS"] = [line for project in content.projects for line in (
            project.title + "  |  " + project.technologies, project.description,
            *["• " + bullet for bullet in project.bullets],
        )]
    else:
        sections["COMPÉTENCES"] = _compact_skills(sections["COMPÉTENCES"])

    document = Document(source)
    body = document._element.body
    for child in list(body):
        if child.tag != W + "sectPr":
            body.remove(child)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.524)
    section.left_margin = section.right_margin = Cm(1.854)
    cols = section._sectPr.find(W + "cols")
    if cols is not None:
        section._sectPr.remove(cols)
    for tag in ("headerReference", "footerReference"):
        for node in list(section._sectPr.findall(W + tag)):
            section._sectPr.remove(node)

    # Measured from the Mews Partners reference. Content stays independent of these roles.
    from docx.enum.style import WD_STYLE_TYPE
    navy, ink, muted, accent = "19354B", "1F2A36", "5C6976", "246477"
    roles = {
        "Name": (20.5, navy, False, 0, 1, 1.0, True),
        "Degree": (9.5, ink, True, 0, 1.4, 1.03, True),
        "Focus": (9, accent, False, 0, 5.3, 1.04, True),
        "Contact": (8, muted, False, 0, 5.5, 1.05, True),
        "Section": (10, navy, True, 6.2, 3.4, 1.0, True),
        "Body": (9, ink, False, 0, 2, 1.1, False),
        "Skill": (9, ink, False, 0, 1.2, 1.08, False),
        "Project": (9, ink, True, 4.2, 1.4, 1.05, True),
        "Project Tech": (8, muted, False, 0, 1.6, 1.05, True),
        "Project Desc": (8.5, muted, False, 0, 1.9, 1.07, True),
        "Bullet": (8.5, ink, False, 0, 1.1, 1.07, False),
        "Entry": (9, ink, True, 2, 1, 1.07, True),
        "Entry Desc": (8.5, muted, False, 0, 1.4, 1.08, False),
        "Last": (9, ink, False, 0, 0, 1.06, False),
    }
    for role, (size, color, bold, before, after, leading, keep) in roles.items():
        name = "CV " + role
        style = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb, style.font.bold = RGBColor.from_string(color), bold
        style.font.italic = False
        f = style.paragraph_format
        f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), leading
        f.keep_with_next, f.keep_together, f.widow_control = keep, True, True
        f.left_indent = f.right_indent = f.first_line_indent = Pt(0)
        f.page_break_before = False
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(9)
    normal.font.color.rgb = RGBColor.from_string(ink)

    def rule(p, color="CAD5DF", size="4", space="3"):
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", size), ("space", space), ("color", color)):
            bottom.set(qn("w:" + key), value)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)

    # Preserve named hyperlinks as well as visible URLs, including body links.
    anchors = document_hyperlinks(source)
    url_pattern = r"https?://[^\s·•]+|(?:www\.)?(?:linkedin\.com/[^\s·•]+|github\.com/[^\s·•]+|[\w-]+\.github\.io[^\s·•]*)|[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}"
    pattern = "|".join([url_pattern] + [re.escape(label) for label in sorted(anchors, key=len, reverse=True)])

    def add_text(p, text, bold=None, color=None):
        cursor = 0
        for match in re.finditer(pattern, text):
            run = p.add_run(text[cursor:match.start()])
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            label = match.group()
            target = anchors.get(label) or (label if label.startswith(("http://", "https://", "mailto:")) else ("mailto:" if "@" in label else "https://") + label)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            value = OxmlElement("w:color"); value.set(qn("w:val"), accent); props.append(value)
            run.append(props)
            value = OxmlElement("w:t"); value.text = label; run.append(value)
            link.append(run); p._p.append(link)
            cursor = match.end()
        run = p.add_run(text[cursor:])
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    def paragraph(text, role="Body"):
        p = document.add_paragraph(style="CV " + role)
        if text.startswith("• "):
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(.36), Cm(-.26)
        add_text(p, text)
        return p

    paragraph(header[0], "Name")
    paragraph(header[1], "Degree")
    paragraph(header[2], "Focus")
    contact = " · ".join(header[3:])
    for text in sections.get("LIENS", []):
        for label in ("Portfolio", "GitHub", "LinkedIn"):
            if text.startswith(label):
                text = text[len(label):].lstrip(" :")
                break
        if text and text not in contact:
            contact += (" · " if contact else "") + text
    for label, target in anchors.items():
        visible = target.removeprefix("https://").removeprefix("http://").removeprefix("mailto:")
        if not any(label in text or visible in text for text in header) and not any(label in text or visible in text for title, values in sections.items() if title != "LIENS" for text in values) and target not in contact and visible not in contact:
            contact += (" · " if contact else "") + target
    p = paragraph(contact, "Contact")
    rule(p, "A1BDCC", "7", "7")
    for title in ATS_SECTIONS:
        values = sections.get(title, [])
        if not values:
            continue
        rule(paragraph(title, "Section"))
        if title in {"LANGUES", "CENTRES D'INTÉRÊT"}:
            paragraph(" · ".join(value.removeprefix("• ") for value in values), "Last")
            continue
        for index, text in enumerate(values):
            if title == "PROJETS":
                if " | " in text:
                    first, rest = re.split(r"\s+\|\s+", text, maxsplit=1)
                    paragraph(first, "Project")
                    paragraph(rest, "Project Tech")
                elif index > 0 and " | " in values[index - 1]:
                    paragraph(text.removeprefix("• "), "Project Desc")
                else:
                    paragraph(text if text.startswith("• ") else "• " + text, "Bullet")
            elif title == "COMPÉTENCES" and " : " in text:
                first, separator, rest = text.partition(" : ")
                p = document.add_paragraph(style="CV Skill")
                add_text(p, first, bold=True, color=navy)
                add_text(p, separator + rest)
            elif title in {"FORMATION", "EXPÉRIENCES"}:
                paragraph(text, "Entry" if " | " in text else "Bullet" if text.startswith("• ") else "Entry Desc")
            else:
                paragraph(text)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".docx", delete=False) as handle:
            temporary = Path(handle.name)
        document.save(temporary)
        validate_docx(temporary)
        validate_cv_layout(temporary)
        if content is None:
            validate_cv_preservation(source, temporary)
        temporary.replace(destination)
    except OSError as error:
        raise CvDocumentError(f"Unable to write resume : {error}") from error
    finally:
        if temporary:
            temporary.unlink(missing_ok=True)


def validate_cv_preservation(source: Path, destination: Path) -> None:
    """Compare semantic text in reading order and all used external hyperlinks."""
    def words(values):
        # Ignore layout separators only; preserve punctuation such as C++, C# and URLs.
        text = " ".join(values)
        text = re.sub(r"\s*(?:[•·]|\s[|:]\s)\s*", " ", text)
        return " ".join(text.split())

    old_header, old = _cv_sections(source)
    new_header, new = _cv_sections(destination)
    if words(old_header[:3]) != words(new_header[:3]) or not words(new_header[3:]).startswith(words(old_header[3:])):
        raise CvDocumentError("Resume identity or contact content changed during conversion")
    for title in ATS_SECTIONS:
        if words(old.get(title, [])) != words(new.get(title, [])):
            raise CvDocumentError(f"Resume content changed during conversion: {title}")

    if not set(document_hyperlinks(source).values()) <= set(document_hyperlinks(destination).values()):
        raise CvDocumentError("Resume hyperlink targets changed during conversion")



def create_cv(source: Path, destination: Path, content: object, replace: bool = False) -> None:
    convert_cv_to_ats(source, destination, content, replace)
